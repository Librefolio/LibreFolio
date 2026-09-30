"""
Test Suite: Danske Bank report sets, phase B — the ``broker_danske_bank`` plugin

Written red-first, before the plugin (plan step 4, §4 B0, "Specifica di dettaglio").
The plugin does not exist yet, so it is looked up in ``BRIMProviderRegistry`` inside
every test (``_danske``): the module always collects, and each test fails on its own,
saying the plugin is missing. The tests whose docstring starts with "Fixture guard"
check the writers of this file, not the product: they pass before and after phase B.

What is pinned (B0, and design v5.3 §3.4.2–§3.4.4, §3.5, §7.1):

- identity and contract: code, name, extensions, priority, version, icon, docs page,
  roles, settlement lag, pre-checkpoint policy, history tag, ``test_sample_sets``;
- recognition: ``can_parse`` (the members and its own combined files; never a foreign
  or unreadable file, and never an exception), ``detect_role``, ``describe_member``
  (coverage per axis, an account fingerprint that never serialises), ``describe_set``;
- reading: a Latin-1, cp1252 or UTF-8 cash statement; the fee header with ``<br>`` or a
  space; the currency column unnamed, named ``Valuutta`` or empty; dates as text or as
  date cells; numbers as cells or as Finnish text;
- classification (S1), rule M (D-S28), pairing (S2–S4), the zones table cell by cell,
  the truth points (C_i, border orphans, A13, proven gaps, the final verification, the
  balance chain, the proofs E1–E4 and their conflicts);
- the combined format: headers, value formats, order, privacy (no ``Säilytystili``), purity;
- the parse of the combined file: transactions, tags, descriptions, assets, todos,
  notices, evidence, and checkpoints and verifications rebuilt from the truth rows;
- the exact facts of the two synthetic sample sets, "main" and "gap".

Edge cases are built in temporary folders by two small writers: an XLSX custody export
(dates as text unless asked otherwise) and a cash statement written the way the bank
writes it (Latin-1, ``;``, LF, newest first, the running ``Saldo`` computed from the
rows). Every value is invented. No server, no database: this module is pure.

Design: LibreFolio_developer_journal/Release_2/Phase_0/26_brimDanskeBank/design-phase00BrimReportSets.md (v5.3)
Plan: LibreFolio_developer_journal/Release_2/Phase_0/26_brimDanskeBank/plan-phase00BrimDanskeBankStep4Implementation.prompt.md, §4 B0
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Set, Tuple

import openpyxl
import pytest

from backend.app.config import PROJECT_ROOT
from backend.app.schemas.brim import COMBINED_REQUIRED_HEADERS, BRIMCombinedTable, BRIMCoverage, BRIMParseOutput, BRIMReportRole, BRIMSetShape, is_fake_asset_id
from backend.app.schemas.common import DateRangeModel
from backend.app.services.brim_provider import BRIMProvider, BRIMSetRequiredError, write_combined_csv
from backend.app.services.provider_registry import BRIMProviderRegistry

# =============================================================================
# CONSTANTS
# =============================================================================

PLUGIN_CODE = "broker_danske_bank"
PLUGIN_MODULE = "backend.app.services.brim_providers.broker_danske_bank"
SAMPLE_DIR = PROJECT_ROOT / "backend" / "app" / "services" / "brim_providers" / "sample_reports"
BROKER_ID = 1
# The deposit account of the custody samples. It is an account number: it never leaves the members.
ACCOUNT = "SYNTH-0001"

MAIN_SET = {"custody": ["danske_bank-custody.xlsx"], "cash": ["danske_bank-cash.csv"]}
GAP_SET = {"custody": ["danske_bank-gap-custody-1.xlsx", "danske_bank-gap-custody-2.xlsx"], "cash": ["danske_bank-gap-cash.csv"]}
SAMPLE_SETS = {"main": MAIN_SET, "gap": GAP_SET}
SAMPLE_ROLES = {name: role for sample_set in SAMPLE_SETS.values() for role, names in sample_set.items() for name in names}

CUSTODY_HEADER: Tuple[Optional[str], ...] = ("Kauppapäivä", "Arvopäivä", "Sijoituskohde", "Määrä", "Kurssi", "Palkkio<br/>sis. Alv", "Summa", None, "Tila", "Toimeksiantotyyppi", "Säilytystili")
CASH_HEADER = ("Pvm", "Saaja/Maksaja", "Määrä", "Saldo", "Tila", "Tarkastus")
LF_HEADERS = ["lf_row_kind", "lf_zone", "lf_reason", "lf_checkpoint", "lf_source", "lf_match_key", "lf_value", "lf_currency", "lf_proof"]
COMBINED_HEADERS = [
    *LF_HEADERS,
    "custody:Kauppapäivä",
    "custody:Arvopäivä",
    "custody:Sijoituskohde",
    "custody:Määrä",
    "custody:Kurssi",
    "custody:Palkkio<br/>sis. Alv",
    "custody:Summa",
    "custody:H",
    "custody:Tila",
    "custody:Toimeksiantotyyppi",
    "cash:Pvm",
    "cash:Saaja/Maksaja",
    "cash:Määrä",
    "cash:Saldo",
    "cash:Tila",
    "cash:Tarkastus",
]

DATA_KINDS = {"pair", "standalone", "excluded", "summarized", "deferred"}
TRUTH_KINDS = {"truth_cash", "truth_position", "verification"}
ZONES = {"before", "window", "gap", "after"}
REASONS = {"no_counterpart", "ambiguous", "status", "unknown_type", "invalid", "not_yet_settled", "outside_cash_coverage", "superseded"}

SOURCE_RE = re.compile(r"^(custody|cash)(?:#([1-9]\d*))?:([1-9]\d*)$")
PROOF_RE = re.compile(r"^(exact:E1|exact:E2|at_least:E3|discarded:E[1-4]:(recent_trades|window_not_covered|excluded_rows|same_day_movement|conflict))$")
MATCH_KEY_RE = re.compile(r"^\d{4}-\d{2}-\d{2}\|-?\d+\.\d{2}\|EUR$")
ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
CASH_VALUE_RE = re.compile(r"^-?\d+\.\d{2}$")
QUANTITY_VALUE_RE = re.compile(r"^-?\d+(\.\d*[1-9])?$")

TAGS = ["import", "danske_bank"]
DEMERGER_TAGS = ["import", "danske_bank", "demerger"]
CHARGES = "danske_trade_charges_included"
DEMERGER = "demerger"
DEMERGER_OLD = "demerger_old_leg"
DANSKE_REASON_CODES = {CHARGES, DEMERGER, DEMERGER_OLD}
# Parse notices and their severity; every excluded reason has its own "excluded_<reason>" warning.
PARSE_NOTICES = {"deferred_rows": "info", "name_mismatch": "warning", "proof_discarded": "info", "balance_chain_broken": "warning", "deposit_assumed": "info"}
SET_NOTICES = {"gap", "overlap_mismatch", "empty_member", "balance_chain_broken", "mixed_accounts_in_file"}
# The todo context keys the wizard reads (developer guide, "context keys the frontend understands").
CHARGES_CONTEXT = ("row", "cash", "currency", "compare_nominal", "split_hint", "split_suggestions")


# =============================================================================
# HELPERS — the plugin is looked up inside each test, never at import time
# =============================================================================


def _missing(what: str) -> None:
    """Fail the current test (never the collection), naming what phase B still has to add."""
    pytest.fail(f"{what}: not implemented yet (BRIM report sets, phase B)", pytrace=False)


def _danske() -> BRIMProvider:
    """The registered Danske Bank plugin."""
    BRIMProviderRegistry.auto_discover()
    plugin = BRIMProviderRegistry.get_provider_instance(PLUGIN_CODE)
    if plugin is None:
        _missing(f"no BRIM plugin is registered as {PLUGIN_CODE} ({PLUGIN_MODULE})")
    return plugin


def _require_contract(name: str) -> None:
    """``BRIMProvider`` itself, not only a subclass, defines ``name``."""
    if not hasattr(BRIMProvider, name):
        _missing(f"BRIMProvider has no {name}")


def _d(text: str) -> date:
    return date.fromisoformat(text)


def _fi_date(day: date) -> str:
    return day.strftime("%d.%m.%Y")


def _from_fi_date(text: str) -> date:
    return datetime.strptime(text, "%d.%m.%Y").date()


def _business_days_after(day: date, count: int) -> date:
    """``day`` plus ``count`` business days: Monday to Friday, no holidays (B0)."""
    while count:
        day += timedelta(days=1)
        if day.weekday() < 5:
            count -= 1
    return day


def _fi_amount(amount: Decimal, *, thousands: str = "") -> str:
    """``amount`` as the bank writes it: comma decimal, no trailing zeros, no plus sign ("-718", "22,5", "-1181,04")."""
    digits = f"{abs(amount):.2f}".rstrip("0").rstrip(".")
    whole, _, cents = digits.partition(".")
    if thousands:
        whole = f"{int(whole):,}".replace(",", thousands)
    return ("-" if amount < 0 else "") + whole + ("," + cents if cents else "")


def _number(text: str) -> Decimal:
    """A number as a combined cell carries it: canonical ("-1181.04") or Finnish text ("−1 181,04", "22,5")."""
    return Decimal(text.replace("\u2212", "-").replace("\u00a0", "").replace(" ", "").replace(",", "."))


def _iso_or_none(text: str) -> Optional[date]:
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _decimal_or_none(text: str) -> Optional[Decimal]:
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


# =============================================================================
# WRITERS — a custody export (XLSX) and a cash statement (CSV), like the bank's
# =============================================================================


def _trade(day: str, name: str, quantity: Any, price: Any, summa: Any, *, value: Optional[str] = None, kind: str = "Rajakurssi", status: str = "Toteutettu", account: str = ACCOUNT, currency: Optional[str] = "EUR") -> List[Any]:
    """One custody trade row; the value date defaults to two business days after the trade date."""
    trade_date = _d(day)
    value_date = _d(value) if value else _business_days_after(trade_date, 2)
    return [trade_date, value_date, name, quantity, price, 0, summa, currency, status, kind, account]


def _tuotto(day: str, name: str, quantity: int, summa: Any, *, account: str = ACCOUNT) -> List[Any]:
    """A ``Tuotto`` (income) row: the shares held, no price, the cash in ``Summa``, value date = trade date."""
    return [_d(day), _d(day), name, quantity, 0, 0, summa, "EUR", "Toteutettu", "Tuotto", account]


def _demerger(day: str, name: str, quantity: int, *, old: bool, summa: Any = None, account: str = ACCOUNT) -> List[Any]:
    """A demerger line, old or new, without cash as the bank writes it (``Summa`` and currency empty)."""
    kind = "Jakautuminen, vanha" if old else "Jakautuminen, uusi"
    return [_d(day), _d(day), name, quantity, 0, 0, summa, None if summa is None else "EUR", "Toteutettu", kind, account]


def _kaamos(day: str, *, account: str = ACCOUNT) -> List[Any]:
    """A one-share purchase (30 + 8 of charges) that pins the first or the last trading day of a segment."""
    return _trade(day, "Kaamos Robotics AB", 1, 30, -38, account=account)


def _xlsx_value(value: Any, date_cells: bool) -> Any:
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day) if date_cells else _fi_date(value)
    return value


def _write_custody(path: Path, rows: Sequence[Any], *, header: Sequence[Optional[str]] = CUSTODY_HEADER, preamble: Sequence[Sequence[Any]] = (), date_cells: bool = False) -> Dict[str, int]:
    """Write a custody export: one sheet, the header, the rows (dates as text unless ``date_cells``).

    Each item of ``rows`` is a row, or ``(key, row)``. Returns ``{key: line}``, lines 1-based.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for extra in preamble:
        sheet.append(list(extra))
    sheet.append(list(header))
    lines: Dict[str, int] = {}
    for item in rows:
        key, row = item if isinstance(item, tuple) else (None, item)
        sheet.append([_xlsx_value(value, date_cells) for value in row])
        if key is not None:
            lines[key] = sheet.max_row
    workbook.save(path)
    return lines


@dataclass(frozen=True)
class _Cash:
    """One cash-statement row, given oldest first.

    ``day`` is ISO, or any other text written as is (an invalid date); ``amount`` is a
    dot-decimal string, or any other text written as is (an invalid amount).
    """

    day: str
    label: str
    amount: str
    key: Optional[str] = None
    status: str = "Toteutunut"
    hidden: bool = False  # booked by the bank, so it moves the balance, but missing from the export


def _write_cash(path: Path, rows: Sequence[_Cash], *, opening: str = "0", encoding: str = "latin-1", thousands: str = "") -> Dict[str, int]:
    """Write a cash statement like the bank's: header, rows newest first, the running ``Saldo`` after each row.

    Only posted rows with a valid date and amount move the balance (``Tila`` =
    ``Toteutunut``, label not ``Varaus``); the others carry it unchanged. So the chain
    of the posted rows always holds, unless a ``hidden`` row breaks it.
    Returns ``{key: line}``, lines 1-based, the header being line 1.
    """
    balance = Decimal(opening)
    rendered: List[Tuple[Optional[str], str]] = []
    for row in rows:
        day = _iso_or_none(row.day)
        amount = _decimal_or_none(row.amount)
        if row.status == "Toteutunut" and not row.label.startswith("Varaus") and day is not None and amount is not None:
            balance += amount
        if row.hidden:
            continue
        cells = [_fi_date(day) if day else row.day, row.label, row.amount if amount is None else _fi_amount(amount, thousands=thousands), _fi_amount(balance, thousands=thousands), row.status, "Ei"]
        rendered.append((row.key, ";".join(cells)))
    rendered.reverse()
    text = "\n".join([";".join(CASH_HEADER), *(line for _key, line in rendered)]) + "\n"
    path.write_bytes(text.encode(encoding))
    return {key: number for number, (key, _line) in enumerate(rendered, start=2) if key is not None}


def _write_raw_custody(path: Path, header: Sequence[Any], rows: Sequence[Sequence[Any]]) -> Path:
    """Write header and rows exactly as given (cell types included)."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in [header, *rows]:
        sheet.append(list(row))
    workbook.save(path)
    return path


def _sample_custody(name: str = "danske_bank-custody.xlsx") -> Tuple[List[Any], List[List[Any]]]:
    """Header and rows of a custody sample, cell values as openpyxl reads them."""
    rows = [list(row) for row in openpyxl.load_workbook(SAMPLE_DIR / name).active.iter_rows(values_only=True)]
    return rows[0], rows[1:]


def _fi_text(value: Any, *, thousands: str, minus: str) -> Any:
    """A numeric cell as Finnish text ("−1 181,04", "41,2", "300"); any other cell unchanged."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return value
    text = str(Decimal(str(value)))
    whole, _, fraction = text.lstrip("-").partition(".")
    whole = f"{int(whole):,}".replace(",", thousands)
    return (minus if text.startswith("-") else "") + whole + ("," + fraction if fraction else "")


class _Builder:
    """The member files of one synthetic set, and where every keyed row landed: ``key -> (role, file number, line)``."""

    def __init__(self, folder: Path) -> None:
        self.folder = folder
        self.custody: List[Path] = []
        self.cash: List[Path] = []
        self.where: Dict[str, Tuple[str, int, int]] = {}
        self.row_counts: Dict[Tuple[str, int], int] = {}

    def add_custody(self, rows: Sequence[Any], **options: Any) -> Path:
        number = len(self.custody) + 1
        path = self.folder / f"custody-{number}.xlsx"
        for key, line in _write_custody(path, rows, **options).items():
            self.where[key] = ("custody", number, line)
        self.custody.append(path)
        self.row_counts[("custody", number)] = len(rows)
        return path

    def add_cash(self, rows: Sequence[_Cash], **options: Any) -> Path:
        number = len(self.cash) + 1
        path = self.folder / f"cash-{number}.csv"
        for key, line in _write_cash(path, rows, **options).items():
            self.where[key] = ("cash", number, line)
        self.cash.append(path)
        self.row_counts[("cash", number)] = sum(1 for row in rows if not row.hidden)
        return path

    def members(self) -> Dict[str, List[Path]]:
        return {"custody": list(self.custody), "cash": list(self.cash)}

    def row(self, records: List[Dict[str, str]], key: str) -> Dict[str, str]:
        role, number, line = self.where[key]
        return _data_row(records, role, number, line)


# =============================================================================
# COMBINED TABLE — reading it back
# =============================================================================


def _records(table: BRIMCombinedTable) -> List[Dict[str, str]]:
    return [dict(zip(table.headers, row, strict=True)) for row in table.rows]


def _sources(record: Dict[str, str]) -> Set[Tuple[str, int, int]]:
    """``lf_source`` as ``{(role, file number, line)}``; a role with a single file is file 1."""
    found = set()
    for token in record["lf_source"].split(" + "):
        match = SOURCE_RE.match(token)
        assert match, f"lf_source {record['lf_source']!r} is not 'role[#n]:line' or two of them joined by ' + '"
        found.add((match.group(1), int(match.group(2) or 1), int(match.group(3))))
    return found


def _data_row(records: List[Dict[str, str]], role: str, number: int, line: int) -> Dict[str, str]:
    """The one data row (not a truth row) built from ``line`` of file ``number`` of ``role``."""
    found = [record for record in records if record["lf_row_kind"] in DATA_KINDS and (role, number, line) in _sources(record)]
    assert len(found) == 1, f"{role} file {number} line {line}: {len(found)} data rows in the combined file {[record['lf_source'] for record in found]}"
    return found[0]


def _truth(records: List[Dict[str, str]], kind: str, as_of: Optional[str] = None) -> List[Dict[str, str]]:
    return [record for record in records if record["lf_row_kind"] == kind and (as_of is None or record["lf_checkpoint"] == as_of)]


def _positions_by_title(records: List[Dict[str, str]], title: str) -> List[Dict[str, str]]:
    return [record for record in _truth(records, "truth_position") if record["custody:Sijoituskohde"] == title]


def _line_count(path: Path) -> int:
    return len(path.read_text(encoding="utf-8-sig").splitlines())


def _cash_type(label: str) -> str:
    """The transaction type of a standalone cash row, by label family (B0, S1)."""
    if label.startswith("Nosto osakesäästötililtä"):
        return "WITHDRAWAL"
    if label.startswith("Vero osakesäästötililtä"):
        return "TAX"
    if label.startswith("Palvelumaksu"):
        return "FEE"
    if label.startswith("Korko"):
        return "INTEREST"
    return "DEPOSIT"


def _owed_transactions(records: List[Dict[str, str]]) -> List[Tuple[Any, ...]]:
    """What the parse owes for each pair and standalone row, in file order:
    ``(type, date, quantity, cash, title, description, tags)``."""
    owed = []
    for record in records:
        kind = record["lf_row_kind"]
        if kind == "pair":
            order = record["custody:Toimeksiantotyyppi"]
            quantity = _number(record["custody:Määrä"])
            tx_type = "DIVIDEND" if order == "Tuotto" else ("BUY" if quantity > 0 else "SELL")
            owed.append((tx_type, _from_fi_date(record["cash:Pvm"]), Decimal(0) if tx_type == "DIVIDEND" else quantity, _number(record["cash:Määrä"]), record["custody:Sijoituskohde"], f"{record['cash:Saaja/Maksaja']} ({order})", TAGS))
        elif kind == "standalone" and record["cash:Pvm"]:
            label = record["cash:Saaja/Maksaja"]
            owed.append((_cash_type(label), _from_fi_date(record["cash:Pvm"]), Decimal(0), _number(record["cash:Määrä"]), None, label, TAGS))
        elif kind == "standalone":
            title = record["custody:Sijoituskohde"]
            owed.append(("ADJUSTMENT", _from_fi_date(record["custody:Arvopäivä"]), _number(record["custody:Määrä"]), None, title, f"{record['custody:Toimeksiantotyyppi']}: {title}", DEMERGER_TAGS))
    return owed


def _assert_each_source_row_once(row_counts: Dict[Tuple[str, int], int], records: List[Dict[str, str]]) -> None:
    """I4, no silent loss: every row of every member lands in exactly one data row of the combined file."""
    seen = Counter(source for record in records if record["lf_row_kind"] in DATA_KINDS for source in _sources(record))
    expected = {(role, number, line) for (role, number), count in row_counts.items() for line in range(2, count + 2)}
    assert sorted(set(seen) - expected) == [], "data rows point at lines the members do not have"
    assert sorted(expected - set(seen)) == [], "member rows missing from the combined file"
    assert sorted(source for source, times in seen.items() if times > 1) == [], "member rows copied into several data rows"


def _sort_key(record: Dict[str, str]) -> Tuple[Any, ...]:
    """B0 order: value date; truth rows at their date, after that day's data rows; then custody before cash, file, line."""
    if record["lf_row_kind"] in TRUTH_KINDS:
        return (_d(record["lf_checkpoint"]), 1)
    day = record["cash:Pvm"] or record["custody:Arvopäivä"]
    role, number, line = min(_sources(record), key=lambda source: (source[0] != "custody", source[1], source[2]))
    return (_from_fi_date(day), 0, 0 if role == "custody" else 1, number, line)


def _expected_tx(tx_type: str, day: str, quantity: str, cash: Optional[str], title: Optional[str], description: str) -> Tuple[Any, ...]:
    """A transaction as ``_transactions`` shows it, from literal strings."""
    return (tx_type, _d(day), Decimal(quantity), None if cash is None else Decimal(cash), title, description, DEMERGER_TAGS if tx_type == "ADJUSTMENT" else TAGS)


# =============================================================================
# PARSE OUTPUT — reading it back
# =============================================================================


def _names(output: BRIMParseOutput) -> Dict[int, Optional[str]]:
    return {asset_id: info.extracted_name for asset_id, info in output.extracted_assets.items()}


def _asset(output: BRIMParseOutput, title: str) -> int:
    ids = [asset_id for asset_id, name in _names(output).items() if name == title]
    assert len(ids) == 1, f"{title!r}: {len(ids)} extracted assets carry this name, one expected"
    return ids[0]


def _transactions(output: BRIMParseOutput) -> List[Tuple[Any, ...]]:
    """``(type, date, quantity, cash, title, description, tags)`` of every transaction, in order."""
    names = _names(output)
    return [(tx.type.value, tx.date, tx.quantity, None if tx.cash is None else tx.cash.amount, None if tx.asset_id is None else names.get(tx.asset_id), tx.description, list(tx.tags or [])) for tx in output.transactions]


def _money(items: Sequence[Any]) -> Dict[str, Decimal]:
    return {item.currency: item.amount for item in items}


def _nonzero(money: Dict[str, Decimal]) -> Dict[str, Decimal]:
    return {code: amount for code, amount in money.items() if amount != 0}


def _checkpoint(output: BRIMParseOutput, as_of: str) -> Any:
    found = [checkpoint for checkpoint in output.checkpoints if checkpoint.as_of == _d(as_of)]
    assert len(found) == 1, f"one checkpoint expected at {as_of}, got {[item.as_of.isoformat() for item in output.checkpoints]}"
    return found[0]


def _positions(output: BRIMParseOutput, checkpoint: Any) -> Dict[str, Tuple[Decimal, str]]:
    names = _names(output)
    return {names.get(position.asset_id): (position.quantity, position.exactness) for position in checkpoint.positions}


def _codes(notices: Sequence[Any]) -> List[str]:
    return [notice.code for notice in notices]


def _essence(output: BRIMParseOutput) -> Dict[str, Any]:
    """What a parse says, without the verbatim text it quotes (evidence, messages, suggestions).

    Two readings of the same data (another encoding, another cell type) must agree on it.
    """
    names = _names(output)
    pinned_context = ("row", "cash", "currency", "compare_nominal", "split_hint", "charges")
    return {
        "transactions": _transactions(output),
        "todos": [(todo.tx_index, todo.field, todo.severity, todo.reason_code, {key: (todo.context or {}).get(key) for key in pinned_context}) for todo in output.field_todos],
        "notices": sorted((notice.code, notice.severity) for notice in output.warnings),
        "checkpoints": [
            (
                checkpoint.as_of,
                checkpoint.kind,
                _money(checkpoint.cash),
                sorted((names.get(position.asset_id), position.quantity, position.exactness) for position in checkpoint.positions),
                checkpoint.absorbed.count,
                _nonzero(_money(checkpoint.absorbed.cash)),
                sorted((row.as_of, row.currency, row.amount, row.label) for row in checkpoint.absorbed.rows),
                _money(checkpoint.absorbed.opening_cash),
            )
            for checkpoint in output.checkpoints
        ],
        "verifications": [(verification.as_of, _money(verification.cash)) for verification in output.verifications],
        "assets": sorted(name or "" for name in names.values()),
    }


# =============================================================================
# SAMPLE SETS AND SCENARIOS — combined once, parsed once
# =============================================================================


def _members(sample_set: Dict[str, List[str]]) -> Dict[str, List[Path]]:
    return {role: [SAMPLE_DIR / name for name in names] for role, names in sample_set.items()}


_SAMPLE_TABLES: Dict[str, BRIMCombinedTable] = {}
_SAMPLE_PARSES: Dict[str, Tuple[Path, BRIMParseOutput]] = {}
_BUILDERS: Dict[str, _Builder] = {}
_SCENARIO_TABLES: Dict[str, BRIMCombinedTable] = {}
_SCENARIO_PARSES: Dict[str, Tuple[Path, BRIMParseOutput]] = {}
SCENARIOS: Dict[str, Callable[[_Builder], None]] = {}


def _sample_table(name: str) -> BRIMCombinedTable:
    """The combined table of sample set ``name`` ("main" or "gap"), combined once per session."""
    plugin = _danske()
    if name not in _SAMPLE_TABLES:
        _SAMPLE_TABLES[name] = plugin.combine(_members(SAMPLE_SETS[name]))
    return _SAMPLE_TABLES[name].model_copy(deep=True)


def _sample_parse(name: str, tmp_path_factory: pytest.TempPathFactory) -> Tuple[BRIMCombinedTable, Path, BRIMParseOutput]:
    """The combined table of a sample set, its combined file (written by the core) and the parse of that file."""
    table = _sample_table(name)
    if name not in _SAMPLE_PARSES:
        path = tmp_path_factory.mktemp(f"danske-{name}") / f"danske_bank-{name}-combined.csv"
        write_combined_csv(path, table)
        _SAMPLE_PARSES[name] = (path, _danske().parse(path, broker_id=BROKER_ID))
    path, output = _SAMPLE_PARSES[name]
    return table, path, output.model_copy(deep=True)


def _scenario_files(name: str, tmp_path_factory: pytest.TempPathFactory) -> _Builder:
    """The member files of a synthetic scenario, written once per session."""
    if name not in _BUILDERS:
        builder = _Builder(tmp_path_factory.mktemp(f"danske-{name}"))
        SCENARIOS[name](builder)
        _BUILDERS[name] = builder
    return _BUILDERS[name]


def _scenario(name: str, tmp_path_factory: pytest.TempPathFactory) -> Tuple[_Builder, BRIMCombinedTable]:
    """The members and the combined table of a synthetic scenario, combined once per session."""
    plugin = _danske()
    builder = _scenario_files(name, tmp_path_factory)
    if name not in _SCENARIO_TABLES:
        _SCENARIO_TABLES[name] = plugin.combine(builder.members())
    return builder, _SCENARIO_TABLES[name].model_copy(deep=True)


def _scenario_parse(name: str, tmp_path_factory: pytest.TempPathFactory) -> Tuple[_Builder, BRIMCombinedTable, Path, BRIMParseOutput]:
    builder, table = _scenario(name, tmp_path_factory)
    if name not in _SCENARIO_PARSES:
        path = builder.folder / "combined.csv"
        write_combined_csv(path, table)
        _SCENARIO_PARSES[name] = (path, _danske().parse(path, broker_id=BROKER_ID))
    path, output = _SCENARIO_PARSES[name]
    return builder, table, path, output.model_copy(deep=True)


def _combined(source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> Tuple[BRIMCombinedTable, Path, BRIMParseOutput]:
    """``("sample", "main")`` or ``("scenario", "zones")``: the combined table, its file and its parse."""
    if source == "sample":
        return _sample_parse(name, tmp_path_factory)
    _builder, table, path, output = _scenario_parse(name, tmp_path_factory)
    return table, path, output


def scenario(name: str) -> Callable[[Callable[[_Builder], None]], Callable[[_Builder], None]]:
    """Register a scenario builder under ``name``."""

    def register(function: Callable[[_Builder], None]) -> Callable[[_Builder], None]:
        SCENARIOS[name] = function
        return function

    return register


# Every scenario starts from a cash balance of 1000 and a deposit on 2021-02-01, and trades
# from 2021-03-01 on: so C1 = 2021-02-28 (a Sunday), the border runs 03-01…03-05, and a
# segment ending 03-31 has its window up to 04-07 (five business days, no holidays).
OPENING = "1000"
DEPOSIT = _Cash("2021-02-01", "Matti Meikäläinen", "500", "deposit")


@scenario("zones")
def _zones(builder: _Builder) -> None:
    """Two segments and a proven gap; the cash statement runs from before the first window to after the last.

    Segment 1: trades 03-01…03-31, C1 = 02-28, border 03-01…03-05, window to 04-07.
    Segment 2: trades 06-01…06-30, C2 = 05-31, border 06-01…06-07, window to 07-07.
    """
    builder.add_custody(
        [
            ("a1", _trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)),  # value 03-03
            ("a2", _trade("2021-03-10", "Lumi Holding Oyj", 5, 10, -58)),  # value 03-12, no cash row
            ("a3", _trade("2021-03-31", "Aurora Metsä Oyj", -4, 25, 92, kind="Päivän kurssi")),  # value 04-02
        ]
    )
    builder.add_custody(
        [
            ("b1", _trade("2021-06-01", "Kaamos Robotics AB", 3, 30, -98)),  # value 06-03
            ("b2", _trade("2021-06-30", "Kaamos Robotics AB", 1, 30, -38)),  # value 07-02
        ]
    )
    builder.add_cash(
        [
            _Cash("2021-01-04", "Matti Meikäläinen", "5000", "c1"),
            _Cash("2021-01-29", "Palvelumaksut 01/2021", "-3.5", "c2"),
            _Cash("2021-02-10", "Osto Revontuli Oyj", "-200", "c3"),
            _Cash("2021-02-15", "Siirto säästötilille", "-50", "c4"),
            _Cash("2021-03-02", "Osto Pohjola Tekniikka O", "-120", "c5"),
            _Cash("2021-03-03", "Osto Aurora Metsä Oyj", "-208", "c6"),
            _Cash("2021-03-17", "Osto Saimaa Energia Oyj", "-300", "c7"),
            _Cash("2021-03-31", "Palvelumaksut 03/2021", "-3.5", "c8"),
            _Cash("2021-04-02", "Myynti Aurora Metsä Oyj", "92", "c9"),
            _Cash("2021-04-06", "Myynti Lumi Palvelut Oyj", "250", "c10"),
            _Cash("2021-05-03", "Matti Meikäläinen", "500", "c11"),
            _Cash("2021-05-10", "Osto Tunturi Pankki Oyj", "-400", "c12"),
            _Cash("2021-06-02", "Osto Saimaa Energia Oyj", "-150", "c13"),
            _Cash("2021-06-03", "Osto Kaamos Robotics AB", "-98", "c14"),
            _Cash("2021-06-30", "Palvelumaksut 06/2021", "-3.5", "c15"),
            _Cash("2021-07-02", "Osto Kaamos Robotics AB", "-38", "c16"),
            _Cash("2021-07-06", "Osto Tunturi Pankki Oyj", "-90", "c17"),
            _Cash("2021-07-15", "Myynti Revontuli Oyj", "300", "c18"),
            _Cash("2021-07-30", "Palvelumaksut 07/2021", "-3.5", "c19"),
        ]
    )


@scenario("merged")
def _merged(builder: _Builder) -> None:
    """Two custody exports with a space between them that no unpaired trade row proves: one segment, 03-01…06-30."""
    builder.add_custody([("a1", _trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)), ("a3", _trade("2021-03-31", "Aurora Metsä Oyj", -4, 25, 92, kind="Päivän kurssi"))])
    builder.add_custody([("b1", _trade("2021-06-01", "Kaamos Robotics AB", 3, 30, -98)), ("b2", _trade("2021-06-30", "Kaamos Robotics AB", 1, 30, -38))])
    builder.add_cash(
        [
            _Cash("2021-01-04", "Matti Meikäläinen", "5000"),
            _Cash("2021-03-03", "Osto Aurora Metsä Oyj", "-208"),
            _Cash("2021-04-02", "Myynti Aurora Metsä Oyj", "92"),
            _Cash("2021-05-03", "Matti Meikäläinen", "500", "deposit"),
            _Cash("2021-05-31", "Palvelumaksut 05/2021", "-3.5", "fee"),
            _Cash("2021-06-03", "Osto Kaamos Robotics AB", "-98"),
            _Cash("2021-07-02", "Osto Kaamos Robotics AB", "-38"),
        ]
    )


@scenario("late")
def _late(builder: _Builder) -> None:
    """A13: the cash statement starts on 03-10, after the custody export (03-01): C1 moves to 03-09.

    The statement opens with the income of a ``Tuotto``: the 30 days before it are all
    before the statement, so its proof cannot be checked (E1 ``window_not_covered``).
    """
    builder.add_custody(
        [
            ("k1", _trade("2021-03-01", "Kaamos Robotics AB", 2, 30, -68)),  # value 03-03, before the statement
            ("t", _tuotto("2021-03-10", "Aurora Metsä Oyj", 100, 15)),
            ("k2", _kaamos("2021-03-31")),  # value 04-02
        ]
    )
    builder.add_cash([_Cash("2021-03-10", "Aurora Metsä 1234567890", "15", "e1"), _Cash("2021-04-02", "Osto Kaamos Robotics AB", "-38", "e2")], opening=OPENING)


@scenario("early_end")
def _early_end(builder: _Builder) -> None:
    """The cash statement ends on 03-25, before the last custody trade (03-31, settled 04-02)."""
    builder.add_custody([("a", _trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)), ("k", _kaamos("2021-03-31"))])
    builder.add_cash([DEPOSIT, _Cash("2021-03-03", "Osto Aurora Metsä Oyj", "-208", "f2"), _Cash("2021-03-25", "Palvelumaksut 03/2021", "-3.5", "f3")], opening=OPENING)


@scenario("hole")
def _hole(builder: _Builder) -> None:
    """Two cash statements that leave 03-06…03-21 uncovered; a trade settled there has no cash to meet."""
    builder.add_custody(
        [
            ("a", _trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)),  # value 03-03
            ("k", _trade("2021-03-10", "Kaamos Robotics AB", 2, 30, -68)),  # value 03-12, in the hole
            ("l", _trade("2021-03-31", "Lumi Holding Oyj", 1, 5, -13)),  # value 04-02
        ]
    )
    builder.add_cash([DEPOSIT, _Cash("2021-03-03", "Osto Aurora Metsä Oyj", "-208", "g2"), _Cash("2021-03-05", "Palvelumaksut 03/2021", "-3.5", "g3")], opening=OPENING)
    builder.add_cash([_Cash("2021-03-22", "Matti Meikäläinen", "100", "h1"), _Cash("2021-04-02", "Osto Lumi Holding Oyj", "-13", "h2"), _Cash("2021-04-09", "Palvelumaksut 04/2021", "-3.5", "h3")], opening="1288.50")


@scenario("chain")
def _chain(builder: _Builder) -> None:
    """A row booked on 03-10 is missing from the statement: the running balance breaks on 03-15."""
    builder.add_custody([("a", _trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)), ("k", _kaamos("2021-03-31"))])
    builder.add_cash(
        [
            DEPOSIT,
            _Cash("2021-03-03", "Osto Aurora Metsä Oyj", "-208"),
            _Cash("2021-03-10", "Palvelumaksut 02/2021", "-40", hidden=True),
            _Cash("2021-03-15", "Palvelumaksut 03/2021", "-3.5", "i3"),
            _Cash("2021-04-02", "Osto Kaamos Robotics AB", "-38"),
        ],
        opening=OPENING,
    )


def _proof_cash(*middle: _Cash, last: str = "2021-05-04") -> List[_Cash]:
    """The cash statement of the proof scenarios: the deposit, the first and the last Kaamos purchase around ``middle``."""
    return [DEPOSIT, _Cash("2021-03-03", "Osto Kaamos Robotics AB", "-38"), *middle, _Cash(last, "Osto Kaamos Robotics AB", "-38")]


@scenario("recent_trade")
def _recent_trade(builder: _Builder) -> None:
    """E1: a purchase of the same title 15 days before the ``Tuotto`` (04-20) makes its quantity no proof."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("a", _trade("2021-04-05", "Aurora Metsä Oyj", 10, 20, -208)), ("t", _tuotto("2021-04-20", "Aurora Metsä Oyj", 60, 9)), ("k2", _kaamos("2021-04-30"))])
    builder.add_cash(_proof_cash(_Cash("2021-04-07", "Osto Aurora Metsä Oyj", "-208"), _Cash("2021-04-20", "Aurora Metsä 1234567890", "9")), opening=OPENING)


@scenario("recent_cash")
def _recent_cash(builder: _Builder) -> None:
    """E1: an unpaired ``Osto`` of the same title 13 days before the ``Tuotto`` makes its quantity no proof."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("t", _tuotto("2021-04-20", "Aurora Metsä Oyj", 60, 9)), ("k2", _kaamos("2021-04-30"))])
    builder.add_cash(_proof_cash(_Cash("2021-04-07", "Osto Aurora Metsä Oyj", "-208", "osto"), _Cash("2021-04-20", "Aurora Metsä 1234567890", "9")), opening=OPENING)


@scenario("same_day")
def _same_day(builder: _Builder) -> None:
    """E2: a sale of the old title on the day of the demerger discards the exact proof; E3 still proves 110."""
    builder.add_custody(
        [
            ("k1", _kaamos("2021-03-01")),
            ("s", _trade("2021-04-01", "Lumi Holding Oyj", -10, 10, 92, kind="Päivän kurssi", value="2021-04-06")),
            ("v", _demerger("2021-04-01", "Lumi Holding Oyj", -100, old=True)),
            ("u", _demerger("2021-04-01", "Lumi Uusi Oyj", 100, old=False)),
            ("k2", _kaamos("2021-04-30")),
        ]
    )
    builder.add_cash(_proof_cash(_Cash("2021-04-06", "Myynti Lumi Holding Oyj", "92")), opening=OPENING)


@scenario("excluded_proof")
def _excluded_proof(builder: _Builder) -> None:
    """E4: a cancelled purchase of the title (03-05) between C1 and the ``Tuotto`` (04-20) discards the proof."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("x", _trade("2021-03-05", "Aurora Metsä Oyj", 10, 20, -208, status="Peruttu")), ("t", _tuotto("2021-04-20", "Aurora Metsä Oyj", 60, 9)), ("k2", _kaamos("2021-04-30"))])
    builder.add_cash(_proof_cash(_Cash("2021-04-20", "Aurora Metsä 1234567890", "9")), opening=OPENING)


@scenario("superseded_proof")
def _superseded_proof(builder: _Builder) -> None:
    """E4 exception: a row superseded by rule M does not discard the proof. The newer export wins 03-10."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("x", _trade("2021-03-10", "Aurora Metsä Oyj", 5, 20, -108))])
    builder.add_custody([("y", _trade("2021-03-10", "Saimaa Energia Oyj", 2, 10, -28)), ("t", _tuotto("2021-04-20", "Aurora Metsä Oyj", 60, 9)), ("k2", _kaamos("2021-04-30"))])
    builder.add_cash(_proof_cash(_Cash("2021-03-12", "Osto Saimaa Energia Oyj", "-28"), _Cash("2021-04-20", "Aurora Metsä 1234567890", "9")), opening=OPENING)


@scenario("conflict")
def _conflict(builder: _Builder) -> None:
    """Two exact proofs that disagree: the ``Tuotto`` says 150, the demerger 120, nothing moved in between."""
    builder.add_custody(
        [
            ("k1", _kaamos("2021-03-01")),
            ("t", _tuotto("2021-04-15", "Lumi Holding Oyj", 150, 10)),
            ("v", _demerger("2021-06-01", "Lumi Holding Oyj", -120, old=True)),
            ("u", _demerger("2021-06-01", "Lumi Uusi Oyj", 120, old=False)),
            ("k2", _kaamos("2021-06-30")),  # value 07-02
        ]
    )
    builder.add_cash(_proof_cash(_Cash("2021-04-15", "Lumi Holding 1234567890", "10"), last="2021-07-02"), opening=OPENING)


@scenario("name_mismatch")
def _name_mismatch(builder: _Builder) -> None:
    """S3: a unique key whose cash label names another title."""
    builder.add_custody([("a", _trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)), ("k2", _kaamos("2021-03-31"))])
    builder.add_cash([DEPOSIT, _Cash("2021-03-03", "Osto Revontuli Oyj", "-208", "c"), _Cash("2021-04-02", "Osto Kaamos Robotics AB", "-38")], opening=OPENING)


@scenario("direction")
def _direction(builder: _Builder) -> None:
    """S2: same date, same amount, but an income against a sale: the direction is part of the key."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("t", _tuotto("2021-03-10", "Aurora Metsä Oyj", 100, 50)), ("k2", _kaamos("2021-03-31"))])
    builder.add_cash(_proof_cash(_Cash("2021-03-10", "Myynti Aurora Metsä Oyj", "50", "c"), last="2021-04-02"), opening=OPENING)


@scenario("by_name")
def _by_name(builder: _Builder) -> None:
    """S4: two titles with the same key; the names tell them apart. The file order alone would cross them."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("a", _trade("2021-03-10", "Aurora Metsä Oyj", 5, 20, -108)), ("s", _trade("2021-03-10", "Saimaa Energia Oyj", 8, 12.5, -108)), ("k2", _kaamos("2021-03-31"))])
    builder.add_cash(_proof_cash(_Cash("2021-03-12", "Osto Aurora Metsä Oyj", "-108", "ca"), _Cash("2021-03-12", "Osto Saimaa Energia Oyj", "-108", "cs"), last="2021-04-02"), opening=OPENING)


@scenario("ambiguous")
def _ambiguous(builder: _Builder) -> None:
    """S4: one truncated cash label compatible with two titles of the same key, and different trades: never guess."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("p1", _trade("2021-03-10", "Pohjola Tekniikka Oyj A", 5, 20, -108)), ("p2", _trade("2021-03-10", "Pohjola Tekniikka Oyj B", 8, 12.5, -108)), ("k2", _kaamos("2021-03-31"))])
    builder.add_cash(_proof_cash(_Cash("2021-03-12", "Osto Pohjola Tekniikka O", "-108", "cp"), last="2021-04-02"), opening=OPENING)


def _crowd(count: int) -> Callable[[_Builder], None]:
    def build(builder: _Builder) -> None:
        fills = [(f"s{index}", _trade("2021-03-10", "Saimaa Energia Oyj", 1, 5, -13)) for index in range(count)]
        builder.add_custody([("k1", _kaamos("2021-03-01")), *fills, ("k2", _kaamos("2021-03-31"))])
        builder.add_cash(_proof_cash(*(_Cash("2021-03-12", "Osto Saimaa Energia Oyj", "-13", f"c{index}") for index in range(count)), last="2021-04-02"), opening=OPENING)

    return build


SCENARIOS["crowd_8"] = _crowd(8)
SCENARIOS["crowd_9"] = _crowd(9)


@scenario("overlap_same")
def _overlap_same(builder: _Builder) -> None:
    """Rule M: two custody exports agree on 03-10, where each holds the same two identical partial fills."""
    fill = _trade("2021-03-10", "Saimaa Energia Oyj", 25, 12.4, -318)
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("a1", list(fill)), ("a2", list(fill))])
    builder.add_custody([("b1", list(fill)), ("b2", list(fill)), ("k2", _kaamos("2021-03-31"))])
    builder.add_cash(_proof_cash(_Cash("2021-03-12", "Osto Saimaa Energia Oyj", "-318", "c1"), _Cash("2021-03-12", "Osto Saimaa Energia Oyj", "-318", "c2"), last="2021-04-02"), opening=OPENING)


def _overlap_diff(newer_first: bool) -> Callable[[_Builder], None]:
    """Rule M: on 03-10 the older export (ends 03-10) says one purchase, the newer one (ends 03-31) two."""

    def build(builder: _Builder) -> None:
        older = [("k1", _kaamos("2021-03-01")), ("x", _trade("2021-03-10", "Aurora Metsä Oyj", 5, 20, -108))]
        newer = [("x2", _trade("2021-03-10", "Aurora Metsä Oyj", 5, 20, -108)), ("y", _trade("2021-03-10", "Saimaa Energia Oyj", 2, 10, -28)), ("k2", _kaamos("2021-03-31"))]
        for rows in [newer, older] if newer_first else [older, newer]:
            builder.add_custody(rows)
        builder.add_cash(_proof_cash(_Cash("2021-03-12", "Osto Aurora Metsä Oyj", "-108", "ca"), _Cash("2021-03-12", "Osto Saimaa Energia Oyj", "-28", "cs"), last="2021-04-02"), opening=OPENING)

    return build


SCENARIOS["overlap_diff"] = _overlap_diff(newer_first=False)
SCENARIOS["overlap_diff_newer_first"] = _overlap_diff(newer_first=True)


@scenario("overlap_cash")
def _overlap_cash(builder: _Builder) -> None:
    """Rule M on the cash role: two statements agree on 03-03…03-12."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("a", _trade("2021-03-10", "Aurora Metsä Oyj", 5, 20, -108)), ("k2", _kaamos("2021-03-31"))])
    builder.add_cash([DEPOSIT, _Cash("2021-03-03", "Osto Kaamos Robotics AB", "-38"), _Cash("2021-03-12", "Osto Aurora Metsä Oyj", "-108")], opening=OPENING)
    builder.add_cash([_Cash("2021-03-03", "Osto Kaamos Robotics AB", "-38"), _Cash("2021-03-12", "Osto Aurora Metsä Oyj", "-108"), _Cash("2021-04-02", "Osto Kaamos Robotics AB", "-38")], opening="1500")


@scenario("mixed_accounts")
def _mixed_accounts(builder: _Builder) -> None:
    """One custody export with two deposit accounts."""
    builder.add_custody([("k1", _kaamos("2021-03-01", account="FI-ACCOUNT-0001")), ("k2", _kaamos("2021-03-31", account="FI-ACCOUNT-0002"))])
    builder.add_cash(_proof_cash(last="2021-04-02"), opening=OPENING)


@scenario("empty_member")
def _empty_member(builder: _Builder) -> None:
    """A second custody export with its header only."""
    builder.add_custody([("k1", _kaamos("2021-03-01")), ("k2", _kaamos("2021-03-31"))])
    builder.add_custody([])
    builder.add_cash(_proof_cash(last="2021-04-02"), opening=OPENING)


@scenario("classes")
def _classes(builder: _Builder) -> None:
    """One row per line of the S1 classes table (B0), in the window of one segment (03-01…03-31)."""
    no_date = _trade("2021-03-22", "Aurora Metsä Oyj", 1, 5, -13, value="2021-03-24")
    no_date[0] = ""
    builder.add_custody(
        [
            ("k0", _kaamos("2021-03-01")),
            ("status", _trade("2021-03-08", "Aurora Metsä Oyj", 5, 20, -108, status="Hylätty")),
            ("buy_positive_summa", _trade("2021-03-09", "Aurora Metsä Oyj", 5, 20, 108)),
            ("sell_negative_summa", _trade("2021-03-09", "Aurora Metsä Oyj", -5, 20, -92, kind="Päivän kurssi")),
            ("pika_buy", _trade("2021-03-11", "Lumi Holding Oyj", 4, 10, -48, kind="Pikakauppa", value="2021-03-15")),
            ("paiva_sell", _trade("2021-03-12", "Kaamos Robotics AB", -1, 33, 25, kind="Päivän kurssi", value="2021-03-16")),
            ("tuotto", _tuotto("2021-03-17", "Aurora Metsä Oyj", 10, 4.5)),
            ("vanha", _demerger("2021-03-18", "Lumi Holding Oyj", -4, old=True)),
            ("uusi", _demerger("2021-03-18", "Lumi Uusi Oyj", 4, old=False)),
            ("uusi_with_summa", _demerger("2021-03-18", "Lumi Toinen Oyj", 4, old=False, summa=-10)),
            ("unknown_type", _trade("2021-03-19", "Aurora Metsä Oyj", 5, 0, None, kind="Siirto", currency=None)),
            ("no_name", _trade("2021-03-22", "", 1, 5, -13, value="2021-03-24")),
            ("no_quantity", _trade("2021-03-22", "Aurora Metsä Oyj", None, 5, -13, value="2021-03-24")),
            ("no_trade_date", no_date),
            ("value_before_trade", _trade("2021-03-23", "Aurora Metsä Oyj", 1, 5, -13, value="2021-03-22")),
            ("k_end", _kaamos("2021-03-31")),
        ]
    )
    builder.add_cash(
        [
            _Cash("2021-02-01", "Matti Meikäläinen", "2000", "deposit_before"),
            _Cash("2021-03-03", "Osto Kaamos Robotics AB", "-38", "p_k0"),
            _Cash("2021-03-04", "Varaus", "-25", "varaus_label"),
            _Cash("2021-03-05", "Osto Revontuli Oyj", "-60", "pending", status="Odottaa"),
            _Cash("2021-03-08", "Nosto osakesäästötililtä", "-100", "nosto"),
            _Cash("2021-03-08", "Vero osakesäästötililtä", "-12.3", "vero"),
            _Cash("2021-03-09", "Palvelumaksut 03/2021", "-3.5", "fee"),
            _Cash("2021-03-10", "Korko 03/2021", "0.42", "korko"),
            _Cash("2021-03-11", "Siirto omalta tililtä", "150", "other_positive"),
            _Cash("2021-03-12", "Siirto ulkomaille", "-40", "other_negative"),
            _Cash("2021-03-12", "Osto Saimaa Energia Oyj", "30", "osto_positive"),
            _Cash("2021-03-15", "Osto Lumi Holding Oyj", "-48", "p_pika"),
            _Cash("2021-03-15", "Myynti Tunturi Pankki Oyj", "-20", "myynti_negative"),
            _Cash("2021-03-16", "Myynti Kaamos Robotics AB", "25", "p_paiva"),
            _Cash("2021-03-16", "Nosto osakesäästötililtä", "10", "nosto_positive"),
            _Cash("2021-03-16", "Palvelumaksut 02/2021", "3.5", "fee_positive"),
            _Cash("2021-03-16", "Korko 02/2021", "-0.1", "korko_negative"),
            _Cash("2021-03-16", "Vero osakesäästötililtä", "5", "vero_positive"),
            _Cash("2021-03-17", "Aurora Metsä 1234567890", "4.5", "p_tuotto"),
            _Cash("2021-03-17", "Revontuli 9876543210", "-7", "income_negative"),
            _Cash("31.02.2021", "Matti Meikäläinen", "10", "bad_date"),
            _Cash("2021-03-18", "Osto Aurora Metsä Oyj", "12,34,5", "bad_amount"),
            _Cash("2021-03-19", "Palvelumaksut 03/2021", "0", "zero_amount"),
            _Cash("2021-03-31", "Palvelumaksut 03/2021", "-3.5", "fee_end"),
            _Cash("2021-04-02", "Osto Kaamos Robotics AB", "-38", "p_k_end"),
        ]
    )


# The expected class of every row of the "classes" scenario: (row kind, reason).
CLASS_TABLE = {
    # custody
    "k0": ("pair", ""),
    "status": ("excluded", "status"),
    "buy_positive_summa": ("excluded", "invalid"),
    "sell_negative_summa": ("excluded", "invalid"),
    "pika_buy": ("pair", ""),
    "paiva_sell": ("pair", ""),
    "tuotto": ("pair", ""),
    "vanha": ("standalone", ""),
    "uusi": ("standalone", ""),
    "uusi_with_summa": ("excluded", "invalid"),
    "unknown_type": ("excluded", "unknown_type"),
    "no_name": ("excluded", "invalid"),
    "no_quantity": ("excluded", "invalid"),
    "no_trade_date": ("excluded", "invalid"),
    "value_before_trade": ("excluded", "invalid"),
    "k_end": ("pair", ""),
    # cash
    "deposit_before": ("standalone", ""),
    "varaus_label": ("excluded", "status"),
    "pending": ("excluded", "status"),
    "nosto": ("standalone", ""),
    "vero": ("standalone", ""),
    "fee": ("standalone", ""),
    "korko": ("standalone", ""),
    "other_positive": ("standalone", ""),
    "other_negative": ("excluded", "unknown_type"),
    "osto_positive": ("excluded", "unknown_type"),
    "myynti_negative": ("excluded", "unknown_type"),
    "nosto_positive": ("excluded", "unknown_type"),
    "fee_positive": ("excluded", "unknown_type"),
    "korko_negative": ("excluded", "unknown_type"),
    "vero_positive": ("excluded", "unknown_type"),
    "income_negative": ("excluded", "unknown_type"),
    "bad_date": ("excluded", "invalid"),
    "bad_amount": ("excluded", "invalid"),
    "zero_amount": ("excluded", "invalid"),
    "fee_end": ("standalone", ""),
}
CLASS_PAIRS = {"k0": "p_k0", "pika_buy": "p_pika", "paiva_sell": "p_paiva", "tuotto": "p_tuotto", "k_end": "p_k_end"}


# =============================================================================
# 1. IDENTITY AND CONTRACT
# =============================================================================


class TestIdentity:
    """B0 "Identità": who the plugin is, the roles it combines, the parameters of its rules."""

    def test_registered_from_its_module(self) -> None:
        plugin = _danske()

        assert plugin.provider_code == PLUGIN_CODE
        assert type(plugin).__module__ == PLUGIN_MODULE

    def test_metadata(self) -> None:
        plugin = _danske()

        assert (plugin.provider_name, plugin.supported_extensions, plugin.detection_priority, plugin.plugin_version) == ("Danske Bank", [".xlsx", ".csv"], 100, "1.0.0")
        assert plugin.description and plugin.description.strip()
        assert (plugin.icon_url, plugin.docs_url) == ("https://danskebank.fi/favicon.ico", "/mkdocs/user/transactions/import/danske-bank/")
        assert plugin.test_file_pattern == "danske_bank"

    def test_report_set_parameters(self) -> None:
        plugin = _danske()

        assert plugin.is_report_set_plugin is True
        assert (plugin.settlement_lag_business_days, plugin.pre_checkpoint_policy, plugin.history_tag) == (5, "summarize", "danske_bank")

    def test_report_roles(self) -> None:
        """custody: XLSX, required, several files, a year deep; cash: CSV, required, several files, five years, covering custody."""
        roles = _danske().report_roles

        assert all(isinstance(role, BRIMReportRole) for role in roles)
        assert [(role.code, role.required, role.multiple, role.extensions, role.max_history, role.must_cover) for role in roles] == [("custody", True, True, [".xlsx"], "P1Y", None), ("cash", True, True, [".csv"], "P5Y", "custody")]
        assert all(role.description.strip() for role in roles)

    def test_plugin_catalogue_shows_the_roles(self) -> None:
        plugin = _danske()

        info = plugin.to_plugin_info()
        catalogue = {item.code: item for item in BRIMProviderRegistry.list_plugin_info()}

        assert (info.code, info.name, info.plugin_version, info.detection_priority, info.icon_url, info.docs_url) == (PLUGIN_CODE, "Danske Bank", "1.0.0", 100, plugin.icon_url, plugin.docs_url)
        assert info.report_roles == plugin.report_roles
        assert [role.code for role in catalogue[PLUGIN_CODE].report_roles] == ["custody", "cash"]

    def test_declares_its_two_sample_sets(self) -> None:
        """The test-only contract property ``test_sample_sets``: one ``{role: [sample names]}`` per set, in this order."""
        plugin = _danske()
        _require_contract("test_sample_sets")

        assert plugin.test_sample_sets == [MAIN_SET, GAP_SET]

    @pytest.mark.parametrize("name", sorted(SAMPLE_ROLES))
    def test_fixture_guard_samples_exist(self, name: str) -> None:
        """Fixture guard: the five synthetic members are in ``sample_reports/`` and carry the test pattern."""
        assert (SAMPLE_DIR / name).is_file()
        assert "danske_bank" in name


# =============================================================================
# 2. RECOGNITION: can_parse and detect_role
# =============================================================================

FOREIGN_SAMPLES = ["credit_agricole-export.csv", "credit_agricole-export.xlsx", "credit_agricole-conti.xlsx", "intesa-lista.xlsx", "intesa-patrimonio.csv", "generic_simple.csv", "degiro-export.csv"]


def _garbage_files() -> List[Any]:
    """Files no plugin can read, and files whose bytes do not match their extension."""
    custody = (SAMPLE_DIR / "danske_bank-custody.xlsx").read_bytes()
    cash = (SAMPLE_DIR / "danske_bank-cash.csv").read_bytes()
    return [
        pytest.param("empty.csv", b"", id="empty-csv"),
        pytest.param("empty.xlsx", b"", id="empty-xlsx"),
        pytest.param("noise.csv", bytes(range(256)) * 4, id="binary-noise-csv"),
        pytest.param("noise.xlsx", bytes(range(256)) * 4, id="binary-noise-xlsx"),
        pytest.param("truncated.xlsx", custody[: len(custody) // 2], id="truncated-xlsx"),
        pytest.param("statement.xlsx", cash, id="cash-text-named-xlsx"),
        pytest.param("custody.csv", custody, id="custody-zip-named-csv"),
    ]


def _write_rows_csv(path: Path, headers: Sequence[str], rows: Sequence[Sequence[str]]) -> Path:
    """A combined-looking file written without the core, so it may lack a column the core would require."""
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";", lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)
    return path


class TestCanParse:
    """B0 "Riconoscimento": the members and its own combined files — nothing else, and never an exception."""

    @pytest.mark.parametrize("name", sorted(SAMPLE_ROLES))
    def test_every_member_sample_is_its_own(self, name: str) -> None:
        plugin = _danske()
        path = SAMPLE_DIR / name

        assert plugin.can_parse(path) is True
        assert PLUGIN_CODE in BRIMProviderRegistry.get_compatible_plugins(path)
        assert BRIMProviderRegistry.auto_detect_plugin(path) == PLUGIN_CODE

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_its_combined_file_is_its_own(self, set_name: str, tmp_path: Path) -> None:
        plugin = _danske()
        path = tmp_path / f"Danske Bank — combined {set_name}.csv"
        write_combined_csv(path, _sample_table(set_name))

        assert plugin.can_parse(path) is True
        assert BRIMProviderRegistry.auto_detect_plugin(path) == PLUGIN_CODE

    @pytest.mark.parametrize("name", FOREIGN_SAMPLES)
    def test_foreign_samples_are_left_to_their_plugins(self, name: str) -> None:
        plugin = _danske()
        path = SAMPLE_DIR / name
        assert path.is_file()  # presence barrier: a missing file is refused too, for another reason

        assert plugin.can_parse(path) is False
        assert plugin.detect_role(path) is None

    @pytest.mark.parametrize(("name", "content"), _garbage_files())
    def test_unreadable_files_are_refused_without_raising(self, name: str, content: bytes, tmp_path: Path) -> None:
        plugin = _danske()
        path = tmp_path / name
        path.write_bytes(content)

        assert plugin.can_parse(path) is False
        assert plugin.detect_role(path) is None

    @pytest.mark.parametrize("name", ["gone.csv", "gone.xlsx"])
    def test_a_missing_file_is_refused_without_raising(self, name: str, tmp_path: Path) -> None:
        assert _danske().can_parse(tmp_path / name) is False

    @pytest.mark.parametrize("header", ["Kauppapäivä", "Arvopäivä", "Sijoituskohde", "Määrä", "Summa", "Toimeksiantotyyppi"])
    def test_a_custody_export_without_a_required_header_is_refused(self, header: str, tmp_path: Path) -> None:
        plugin = _danske()
        rows = [_trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)]
        _write_custody(tmp_path / "control.xlsx", rows)
        _write_custody(tmp_path / "custody.xlsx", rows, header=["Muu sarake" if cell == header else cell for cell in CUSTODY_HEADER])

        assert plugin.can_parse(tmp_path / "control.xlsx") is True  # presence barrier: the same file with every header
        assert plugin.can_parse(tmp_path / "custody.xlsx") is False

    @pytest.mark.parametrize("header", ["Pvm", "Saaja/Maksaja", "Määrä", "Saldo", "Tila"])
    def test_a_cash_statement_without_a_required_header_is_refused(self, header: str, tmp_path: Path) -> None:
        plugin = _danske()
        _write_cash(tmp_path / "control.csv", [DEPOSIT])
        path = tmp_path / "cash.csv"
        _write_cash(path, [DEPOSIT])
        lines = path.read_bytes().decode("latin-1").split("\n")
        lines[0] = ";".join("Muu" if cell == header else cell for cell in CASH_HEADER)
        path.write_bytes("\n".join(lines).encode("latin-1"))

        assert plugin.can_parse(tmp_path / "control.csv") is True  # presence barrier
        assert plugin.can_parse(path) is False

    def test_the_custody_header_is_found_below_a_preamble(self, tmp_path: Path) -> None:
        """B0: the custody headers are looked for down to row 20; here they are on row 4."""
        plugin = _danske()
        path = tmp_path / "custody.xlsx"
        _write_custody(path, [_trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208)], preamble=[["Sijoitukset"], ["Tapahtumat 01.03.2021 - 31.03.2021"], [None]])

        assert plugin.can_parse(path) is True
        assert plugin.detect_role(path) == "custody"

    def test_another_plugins_combined_file_is_not_its_own(self, tmp_path: Path) -> None:
        """``lf_`` columns alone do not make a Danske combined file: its own role columns must be there too."""
        plugin = _danske()
        path = _write_rows_csv(tmp_path / "combined.csv", ["lf_row_kind", "lf_source", "custody:trade_date", "custody:asset", "cash:value_date", "cash:amount"], [["standalone", "custody:2", "2025-01-02", "ACME Oyj", "", ""]])

        assert plugin.can_parse(path) is False

    @pytest.mark.parametrize("dropped", ["lf_row_kind", "lf_source", "custody:Toimeksiantotyyppi", "cash:Saaja/Maksaja"])
    def test_its_combined_file_without_a_marker_column_is_refused(self, dropped: str, tmp_path: Path) -> None:
        plugin = _danske()
        table = _sample_table("main")
        keep = [index for index, header in enumerate(table.headers) if header != dropped]
        control = _write_rows_csv(tmp_path / "control.csv", table.headers, table.rows)
        path = _write_rows_csv(tmp_path / "combined.csv", [table.headers[index] for index in keep], [[row[index] for index in keep] for row in table.rows])

        assert plugin.can_parse(control) is True  # presence barrier: the very same file with the column
        assert plugin.can_parse(path) is False


class TestDetectRole:
    """B0: ``custody``, ``cash``, or None — None for its own combined file too."""

    @pytest.mark.parametrize("name", sorted(SAMPLE_ROLES))
    def test_member_samples(self, name: str) -> None:
        assert _danske().detect_role(SAMPLE_DIR / name) == SAMPLE_ROLES[name]

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_a_combined_file_has_no_role(self, set_name: str, tmp_path: Path) -> None:
        plugin = _danske()
        path = tmp_path / "combined.csv"
        write_combined_csv(path, _sample_table(set_name))

        assert plugin.can_parse(path) is True  # presence barrier: it is the plugin's file, only not a member
        assert plugin.detect_role(path) is None


# =============================================================================
# 3. DESCRIBING MEMBERS AND SETS, without combining
# =============================================================================

MEMBER_FACTS = [
    pytest.param("danske_bank-custody.xlsx", "custody", 15, "trade", "2020-02-03", "2020-06-26", id="main-custody"),
    pytest.param("danske_bank-cash.csv", "cash", 32, "value", "2019-01-07", "2020-07-03", id="main-cash"),
    pytest.param("danske_bank-gap-custody-1.xlsx", "custody", 3, "trade", "2020-09-01", "2020-10-30", id="gap-custody-1"),
    pytest.param("danske_bank-gap-custody-2.xlsx", "custody", 4, "trade", "2021-01-05", "2021-02-26", id="gap-custody-2"),
    pytest.param("danske_bank-gap-cash.csv", "cash", 14, "value", "2020-08-03", "2021-03-01", id="gap-cash"),
]


class TestDescribeMember:
    """B0: data rows, coverage on the role's axis from the valid dates, and an account fingerprint kept in memory."""

    @pytest.mark.parametrize(("name", "role", "rows", "axis", "start", "end"), MEMBER_FACTS)
    def test_samples(self, name: str, role: str, rows: int, axis: str, start: str, end: str) -> None:
        summary = _danske().describe_member(SAMPLE_DIR / name)

        assert (summary.role, summary.rows) == (role, rows)
        assert summary.coverage == [BRIMCoverage(axis=axis, start=_d(start), end=_d(end))]

    def test_the_custody_fingerprint_stays_in_memory(self) -> None:
        summary = _danske().describe_member(SAMPLE_DIR / "danske_bank-custody.xlsx")

        assert summary.account_fingerprint, "the deposit check needs a fingerprint of the custody account"
        assert ACCOUNT not in summary.account_fingerprint, "a digest of the account, not the account"
        serialised = json.dumps(summary.model_dump(mode="json")) + summary.model_dump_json()
        assert "account_fingerprint" not in serialised
        assert summary.account_fingerprint not in serialised
        assert ACCOUNT not in serialised

    def test_a_cash_statement_has_no_fingerprint(self) -> None:
        """If it had one, the framework would compare it with the custody one and always report mixed accounts."""
        assert _danske().describe_member(SAMPLE_DIR / "danske_bank-cash.csv").account_fingerprint is None

    def test_fingerprints_tell_the_accounts_apart(self, tmp_path: Path) -> None:
        plugin = _danske()
        prints = {}
        for name, account in (("first", "FI-ACCOUNT-0001"), ("second", "FI-ACCOUNT-0002"), ("first_again", "FI-ACCOUNT-0001")):
            path = tmp_path / f"{name}.xlsx"
            _write_custody(path, [_trade("2021-03-01", "Aurora Metsä Oyj", 10, 20, -208, account=account)])
            prints[name] = plugin.describe_member(path).account_fingerprint

        assert all(prints.values())
        assert prints["first"] != prints["second"]
        assert prints["first"] == prints["first_again"]

    def test_coverage_counts_valid_dates_only(self, tmp_path: Path) -> None:
        plugin = _danske()
        broken = _trade("2021-03-15", "Aurora Metsä Oyj", 1, 5, -13)
        broken[0] = "31.02.2021"
        path = tmp_path / "custody.xlsx"
        _write_custody(path, [_kaamos("2021-03-01"), broken, _kaamos("2021-03-31")])

        assert plugin.describe_member(path).coverage == [BRIMCoverage(axis="trade", start=_d("2021-03-01"), end=_d("2021-03-31"))]

    @pytest.mark.parametrize("role", ["custody", "cash"])
    def test_an_export_with_its_header_only(self, role: str, tmp_path: Path) -> None:
        """Design §3.3: role recognised, zero rows, no coverage."""
        plugin = _danske()
        path = tmp_path / ("custody.xlsx" if role == "custody" else "cash.csv")
        if role == "custody":
            _write_custody(path, [])
        else:
            _write_cash(path, [])

        summary = plugin.describe_member(path)

        assert (plugin.detect_role(path), summary.role, summary.rows, summary.coverage) == (role, role, 0, [])


class TestDescribeSet:
    """B0: the final merged segments (trade axis), the proven gaps ``[E_i + 1, S_i+1 − 1]``, and notices with stable codes."""

    def test_main_set(self) -> None:
        shape = _danske().describe_set(_members(MAIN_SET))

        assert isinstance(shape, BRIMSetShape)
        assert (shape.segments, shape.gaps) == ([DateRangeModel(start=_d("2020-02-03"), end=_d("2020-06-26"))], [])
        assert _codes(shape.notices) == []

    def test_gap_set(self) -> None:
        shape = _danske().describe_set(_members(GAP_SET))

        assert shape.segments == [DateRangeModel(start=_d("2020-09-01"), end=_d("2020-10-30")), DateRangeModel(start=_d("2021-01-05"), end=_d("2021-02-26"))]
        assert shape.gaps == [DateRangeModel(start=_d("2020-10-31"), end=_d("2021-01-04"))]
        assert _codes(shape.notices) == ["gap"]
        context = shape.notices[0].context or {}
        assert {key: context.get(key) for key in ("start", "end")} == {"start": "2020-10-31", "end": "2021-01-04"}

    def test_a_proven_gap_keeps_the_segments_apart(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        shape = _danske().describe_set(_scenario_files("zones", tmp_path_factory).members())

        assert shape.segments == [DateRangeModel(start=_d("2021-03-01"), end=_d("2021-03-31")), DateRangeModel(start=_d("2021-06-01"), end=_d("2021-06-30"))]
        assert shape.gaps == [DateRangeModel(start=_d("2021-04-01"), end=_d("2021-05-31"))]
        assert _codes(shape.notices).count("gap") == 1

    def test_an_unproven_gap_merges_the_segments(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """No unpaired trade row of the cash statement falls between the two custody exports: one segment."""
        shape = _danske().describe_set(_scenario_files("merged", tmp_path_factory).members())

        assert (shape.segments, shape.gaps) == ([DateRangeModel(start=_d("2021-03-01"), end=_d("2021-06-30"))], [])
        assert "gap" not in _codes(shape.notices)

    def test_a_late_cash_statement_does_not_move_the_segment(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """A13 moves the checkpoint, not the segment: the segment is the custody coverage."""
        shape = _danske().describe_set(_scenario_files("late", tmp_path_factory).members())

        assert shape.segments == [DateRangeModel(start=_d("2021-03-01"), end=_d("2021-03-31"))]

    @pytest.mark.parametrize(("name", "flagged"), [("overlap_diff", True), ("overlap_diff_newer_first", True), ("overlap_same", False), ("overlap_cash", False)])
    def test_overlap_mismatch(self, name: str, flagged: bool, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Rule M: files of one role that disagree on a day they both cover are reported; files that agree are not."""
        shape = _danske().describe_set(_scenario_files(name, tmp_path_factory).members())

        assert ("overlap_mismatch" in _codes(shape.notices)) is flagged
        assert shape.segments == [DateRangeModel(start=_d("2021-03-01"), end=_d("2021-03-31"))]

    def test_an_empty_member(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        plugin = _danske()
        builder = _scenario_files("empty_member", tmp_path_factory)

        shape = plugin.describe_set(builder.members())

        assert "empty_member" in _codes(shape.notices)
        assert shape.segments == [DateRangeModel(start=_d("2021-03-01"), end=_d("2021-03-31"))]
        assert plugin.describe_member(builder.custody[1]).rows == 0

    def test_a_broken_balance_chain(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        shape = _danske().describe_set(_scenario_files("chain", tmp_path_factory).members())

        assert "balance_chain_broken" in _codes(shape.notices)

    def test_two_accounts_in_one_custody_export(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        shape = _danske().describe_set(_scenario_files("mixed_accounts", tmp_path_factory).members())

        assert "mixed_accounts_in_file" in _codes(shape.notices)
        dumped = shape.model_dump_json()
        assert "FI-ACCOUNT-0001" not in dumped and "FI-ACCOUNT-0002" not in dumped

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_never_reveals_the_account(self, set_name: str) -> None:
        shape = _danske().describe_set(_members(SAMPLE_SETS[set_name]))

        assert set(_codes(shape.notices)) <= SET_NOTICES
        assert ACCOUNT not in shape.model_dump_json()


# =============================================================================
# 4. READING: encodings, header spellings, cell types
# =============================================================================


def _main_cash_text() -> str:
    return (SAMPLE_DIR / "danske_bank-cash.csv").read_bytes().decode("latin-1")


def _rename_header(position: int, name: str) -> Callable[[List[Any], List[List[Any]]], None]:
    def rewrite(header: List[Any], _rows: List[List[Any]]) -> None:
        header[position] = name

    return rewrite


def _rewrite_cells(first: int, last: int, cell: Callable[[Any], Any]) -> Callable[[List[Any], List[List[Any]]], None]:
    def rewrite(_header: List[Any], rows: List[List[Any]]) -> None:
        for row in rows:
            row[first:last] = [cell(value) for value in row[first:last]]

    return rewrite


# The ways a custody export may differ from the sample without being different data.
CUSTODY_VARIANTS: Dict[str, Callable[[List[Any], List[List[Any]]], None]] = {
    "br": _rename_header(5, "Palkkio<br>sis. Alv"),
    "space": _rename_header(5, "Palkkio sis. Alv"),
    "valuutta": _rename_header(7, "Valuutta"),
    "no_currency": _rewrite_cells(7, 8, lambda _value: None),
    "date_cells": _rewrite_cells(0, 2, lambda value: datetime.strptime(value, "%d.%m.%Y")),
    "text_numbers_nbsp": _rewrite_cells(3, 7, lambda value: _fi_text(value, thousands="\u00a0", minus="\u2212")),
    "text_numbers_space": _rewrite_cells(3, 7, lambda value: _fi_text(value, thousands=" ", minus="-")),
}


def _custody_variant(variant: str) -> Tuple[List[Any], List[List[Any]]]:
    """The main custody sample, rewritten one way: a header spelling, a missing currency, date cells, Finnish text numbers."""
    header, rows = _sample_custody()
    CUSTODY_VARIANTS[variant](header, rows)
    return header, rows


def _combine_and_parse(custody: Path, cash: Path, folder: Path) -> Tuple[BRIMCombinedTable, BRIMParseOutput]:
    plugin = _danske()
    table = plugin.combine({"custody": [custody], "cash": [cash]})
    path = folder / "combined.csv"
    write_combined_csv(path, table)
    return table, plugin.parse(path, broker_id=BROKER_ID)


def _lf_columns(table: BRIMCombinedTable) -> List[List[str]]:
    return [row[: len(LF_HEADERS)] for row in table.rows]


class TestReadingRobustness:
    """B0 "Lettura": the same data from another encoding, header spelling or cell type makes the same import."""

    @pytest.mark.parametrize("encoding", ["cp1252", "utf-8", "utf-8-sig"])
    def test_cash_statement_encodings(self, encoding: str, tmp_path: Path, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Latin-1 (the sample), cp1252, UTF-8 with and without BOM: the same combined table, the same parse."""
        base_table, _path, base = _sample_parse("main", tmp_path_factory)
        cash = tmp_path / "danske_bank-cash.csv"
        cash.write_bytes(_main_cash_text().encode(encoding))

        table, output = _combine_and_parse(SAMPLE_DIR / "danske_bank-custody.xlsx", cash, tmp_path)

        assert (table.headers, table.rows) == (base_table.headers, base_table.rows)
        assert _essence(output) == _essence(base)

    @pytest.mark.parametrize("separator", [" ", "\u00a0"], ids=["space", "nbsp"])
    def test_cash_amounts_with_thousands_separators(self, separator: str, tmp_path: Path, tmp_path_factory: pytest.TempPathFactory) -> None:
        """F4: Finnish amounts may group thousands with a space or a no-break space; they are the same amounts."""
        base_table, _path, base = _sample_parse("main", tmp_path_factory)
        lines = _main_cash_text().split("\n")
        rewritten = [lines[0]]
        for line in lines[1:]:
            cells = line.split(";")
            if len(cells) == len(CASH_HEADER):
                cells[2], cells[3] = _fi_amount(_number(cells[2]), thousands=separator), _fi_amount(_number(cells[3]), thousands=separator)
            rewritten.append(";".join(cells))
        cash = tmp_path / "danske_bank-cash.csv"
        cash.write_bytes("\n".join(rewritten).encode("latin-1"))
        assert f"6{separator}000" in cash.read_bytes().decode("latin-1")  # presence barrier: the grouping is really there

        table, output = _combine_and_parse(SAMPLE_DIR / "danske_bank-custody.xlsx", cash, tmp_path)

        assert _lf_columns(table) == _lf_columns(base_table)
        assert _essence(output) == _essence(base)

    @pytest.mark.parametrize("variant", ["br", "space", "valuutta", "date_cells"])
    def test_same_cells_under_another_header_or_cell_type(self, variant: str, tmp_path: Path, tmp_path_factory: pytest.TempPathFactory) -> None:
        """``<br>`` or a space in the fee header, the currency column named ``Valuutta``, dates as date cells: the same rows."""
        base_table, _path, base = _sample_parse("main", tmp_path_factory)
        custody = _write_raw_custody(tmp_path / "danske_bank-custody.xlsx", *_custody_variant(variant))

        table, output = _combine_and_parse(custody, SAMPLE_DIR / "danske_bank-cash.csv", tmp_path)

        assert table.rows == base_table.rows
        assert _essence(output) == _essence(base)

    @pytest.mark.parametrize("variant", ["no_currency", "text_numbers_nbsp", "text_numbers_space"])
    def test_same_import_from_other_cell_values(self, variant: str, tmp_path: Path, tmp_path_factory: pytest.TempPathFactory) -> None:
        """An empty currency column means EUR; Finnish text numbers ("−1 181,04") are numbers: the lf_ columns and the parse stay."""
        base_table, _path, base = _sample_parse("main", tmp_path_factory)
        custody = _write_raw_custody(tmp_path / "danske_bank-custody.xlsx", *_custody_variant(variant))

        table, output = _combine_and_parse(custody, SAMPLE_DIR / "danske_bank-cash.csv", tmp_path)

        assert _lf_columns(table) == _lf_columns(base_table)
        assert _essence(output) == _essence(base)

    def test_text_cells_are_copied_as_they_are(self, tmp_path: Path) -> None:
        """Verbatim: a number written as text stays that text in the combined file."""
        custody = _write_raw_custody(tmp_path / "danske_bank-custody.xlsx", *_custody_variant("text_numbers_nbsp"))

        table = _danske().combine({"custody": [custody], "cash": [SAMPLE_DIR / "danske_bank-cash.csv"]})

        kaamos = _data_row(_records(table), "custody", 1, 5)
        assert (kaamos["custody:Määrä"], kaamos["custody:Kurssi"], kaamos["custody:Summa"]) == ("300", "41,2", "\u22121\u00a0181,04")


# =============================================================================
# 5. CLASSES (S1)
# =============================================================================


class TestClassification:
    """S1: one row for every line of the B0 classes table, all in the window of one segment."""

    @pytest.mark.parametrize("key", sorted(CLASS_TABLE))
    def test_row_class(self, key: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table = _scenario("classes", tmp_path_factory)

        row = builder.row(_records(table), key)

        assert (row["lf_row_kind"], row["lf_reason"]) == CLASS_TABLE[key]

    @pytest.mark.parametrize(("custody", "cash"), sorted(CLASS_PAIRS.items()))
    def test_pairable_rows_meet(self, custody: str, cash: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """A purchase (``Rajakurssi``, ``Pikakauppa``), a sale (``Päivän kurssi``) and an income (``Tuotto``) each meet their cash row."""
        builder, table = _scenario("classes", tmp_path_factory)
        records = _records(table)

        row = builder.row(records, custody)

        assert row["lf_row_kind"] == "pair"
        assert row is builder.row(records, cash)

    def test_every_member_row_lands_in_exactly_one_data_row(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table = _scenario("classes", tmp_path_factory)

        _assert_each_source_row_once(builder.row_counts, _records(table))

    def test_standalone_cash_rows_become_their_family_type(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """``Nosto`` → WITHDRAWAL, ``Vero`` → TAX, ``Palvelumaksu…`` → FEE, ``Korko…`` → INTEREST, any other income → DEPOSIT."""
        _builder, _table, _path, output = _scenario_parse("classes", tmp_path_factory)
        expected = {
            ("2021-02-01", "Matti Meikäläinen"): ("DEPOSIT", "2000"),
            ("2021-03-08", "Nosto osakesäästötililtä"): ("WITHDRAWAL", "-100"),
            ("2021-03-08", "Vero osakesäästötililtä"): ("TAX", "-12.3"),
            ("2021-03-09", "Palvelumaksut 03/2021"): ("FEE", "-3.5"),
            ("2021-03-10", "Korko 03/2021"): ("INTEREST", "0.42"),
            ("2021-03-11", "Siirto omalta tililtä"): ("DEPOSIT", "150"),
            ("2021-03-31", "Palvelumaksut 03/2021"): ("FEE", "-3.5"),
        }

        found = {(tx.date.isoformat(), tx.description): tx for tx in output.transactions if tx.asset_id is None}

        assert sorted(found) == sorted(expected)
        for key, (tx_type, amount) in expected.items():
            tx = found[key]
            assert (tx.type.value, tx.quantity, tx.cash.code, tx.cash.amount, tx.tags) == (tx_type, Decimal(0), "EUR", Decimal(amount), TAGS), key

    def test_an_assumed_deposit_is_explained(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """A credit with a label of no known family is booked as a deposit (only the holder pays into the account), with an info notice."""
        _builder, _table, _path, output = _scenario_parse("classes", tmp_path_factory)

        notices = [notice for notice in output.warnings if notice.code == "deposit_assumed"]

        assert notices and all(notice.severity == "info" for notice in notices)


# =============================================================================
# 6. RULE M (D-S28): several files of one role
# =============================================================================


class TestRuleM:
    """Per role, the file that ends last wins each day it covers; identical days keep one copy, differing days exclude the loser's rows."""

    def test_identical_overlap_keeps_one_copy_of_each_row(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Both custody exports hold the same two partial fills on 03-10: two pairs, not four, not one."""
        builder, table = _scenario("overlap_same", tmp_path_factory)
        records = _records(table)

        saimaa = [record for record in records if record["lf_row_kind"] in DATA_KINDS and record["custody:Sijoituskohde"] == "Saimaa Energia Oyj"]

        assert [record["lf_row_kind"] for record in saimaa] == ["pair", "pair"]
        assert {builder.row(records, "c1")["lf_source"], builder.row(records, "c2")["lf_source"]} == {record["lf_source"] for record in saimaa}
        assert [record["lf_source"] for record in records if record["lf_row_kind"] == "excluded"] == []

    @pytest.mark.parametrize(("name", "loser_file"), [("overlap_diff", 1), ("overlap_diff_newer_first", 2)])
    def test_differing_overlap_the_newer_file_wins(self, name: str, loser_file: int, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The export ending 03-31 wins 03-10 over the one ending 03-10, whatever the upload order; the loser's row is ``superseded``."""
        builder, table = _scenario(name, tmp_path_factory)
        records = _records(table)
        assert builder.where["x"][1] == loser_file  # fixture guard: the older export is where the parametrization says

        loser = builder.row(records, "x")

        assert (loser["lf_row_kind"], loser["lf_reason"]) == ("excluded", "superseded")
        assert [record["lf_source"] for record in records if record["lf_reason"] == "superseded"] == [loser["lf_source"]]
        for custody, cash in (("x2", "ca"), ("y", "cs"), ("k1", None)):
            row = builder.row(records, custody)
            assert row["lf_row_kind"] == "pair", custody
            if cash is not None:
                assert row is builder.row(records, cash)

    def test_identical_overlap_of_cash_statements(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _builder, table = _scenario("overlap_cash", tmp_path_factory)

        kinds = Counter(record["lf_row_kind"] for record in _records(table) if record["lf_row_kind"] in DATA_KINDS)

        assert kinds == Counter({"pair": 3, "standalone": 1})


# =============================================================================
# 7. PAIRING (S2–S4)
# =============================================================================


class TestPairing:
    """The key is value date, amount to the cent, currency and direction; the name is a check; one to one; never a guess."""

    def test_the_direction_is_part_of_the_key(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """An income and a sale of the same amount on the same day do not meet: each stays without counterpart."""
        builder, table = _scenario("direction", tmp_path_factory)
        records = _records(table)

        for key in ("t", "c"):
            row = builder.row(records, key)
            assert (row["lf_row_kind"], row["lf_reason"]) == ("excluded", "no_counterpart"), key

    def test_a_unique_key_pairs_despite_another_name(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """S3: the pair is made, with a ``name_mismatch`` warning; the asset comes from the custody row."""
        builder, table, _path, output = _scenario_parse("name_mismatch", tmp_path_factory)
        records = _records(table)

        pair = builder.row(records, "a")

        assert pair["lf_row_kind"] == "pair" and pair is builder.row(records, "c")
        notices = [notice for notice in output.warnings if notice.code == "name_mismatch"]
        assert notices and all(notice.severity == "warning" for notice in notices)
        [buy] = [tx for tx in output.transactions if tx.type.value == "BUY" and tx.date == _d("2021-03-03")]
        assert (_names(output)[buy.asset_id], buy.description) == ("Aurora Metsä Oyj", "Osto Revontuli Oyj (Rajakurssi)")

    def test_names_sort_out_two_titles_with_the_same_key(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """S4: among the matchings of maximum size, the one with most compatible names — here not the file order."""
        builder, table, _path, output = _scenario_parse("by_name", tmp_path_factory)
        records = _records(table)

        assert builder.row(records, "a") is builder.row(records, "ca")
        assert builder.row(records, "s") is builder.row(records, "cs")
        assert "name_mismatch" not in _codes(output.warnings)

    def test_candidates_that_cannot_be_told_apart_are_all_excluded(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """A truncated label compatible with two titles, and two different trades: every row of the key is ``ambiguous``."""
        builder, table = _scenario("ambiguous", tmp_path_factory)
        records = _records(table)

        for key in ("p1", "p2", "cp"):
            row = builder.row(records, key)
            assert (row["lf_row_kind"], row["lf_reason"]) == ("excluded", "ambiguous"), key

    def test_identical_partial_fills_are_two_pairs(self) -> None:
        """The main sample buys Saimaa Energia twice on 2020-03-25, 25 × 12.4 each: two pairs, one cash row each."""
        records = _records(_sample_table("main"))

        fills = [_data_row(records, "custody", 1, line) for line in (14, 15)]

        assert [record["lf_row_kind"] for record in fills] == ["pair", "pair"]
        assert fills[0] is not fills[1]
        assert sorted(line for record in fills for role, _number, line in _sources(record) if role == "cash") == [13, 14]

    @pytest.mark.parametrize("count", [8, 9])
    def test_more_than_eight_candidates_a_side_are_ambiguous(self, count: int, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Identical fills pair up to eight a side; beyond that the search stops and every row of the key is ``ambiguous``."""
        builder, table = _scenario(f"crowd_{count}", tmp_path_factory)
        records = _records(table)

        rows = [builder.row(records, f"s{index}") for index in range(count)] + [builder.row(records, f"c{index}") for index in range(count)]

        if count <= 8:
            assert all(row["lf_row_kind"] == "pair" for row in rows)
            assert len({row["lf_source"] for row in rows}) == count
        else:
            assert all((row["lf_row_kind"], row["lf_reason"]) == ("excluded", "ambiguous") for row in rows)


# =============================================================================
# 8. ZONES AND OUTCOMES (design §3.4.2, B0 table)
# =============================================================================

# key, row kind, accepted zones, reason, checkpoint. Where B0 says only that a row "follows
# the next zone" (a trade settled after the end of a custody export, inside its window),
# the outcome is pinned and the zone may be the physical one or the next one.
ZONE_CELLS = [
    pytest.param("c1", "standalone", {"before"}, "", "2021-02-28", id="before-standalone-absorbed"),
    pytest.param("c2", "standalone", {"before"}, "", "2021-02-28", id="before-fee-absorbed"),
    pytest.param("c3", "summarized", {"before"}, "", "2021-02-28", id="before-unpaired-trade-summarized"),
    pytest.param("c4", "excluded", {"before"}, "unknown_type", "2021-02-28", id="before-excluded-still-absorbed"),
    pytest.param("c5", "summarized", {"window"}, "", "2021-02-28", id="border-orphan-segment-1"),
    pytest.param("a1", "pair", {"window"}, "", "", id="window-pair"),
    pytest.param("a2", "excluded", {"window"}, "no_counterpart", "", id="window-trade-without-cash"),
    pytest.param("c7", "excluded", {"window"}, "no_counterpart", "", id="window-cash-without-trade"),
    pytest.param("c8", "standalone", {"window"}, "", "", id="window-standalone"),
    pytest.param("c10", "summarized", {"window", "gap"}, "", "2021-05-31", id="beyond-end-follows-the-gap"),
    pytest.param("c11", "standalone", {"gap"}, "", "", id="gap-standalone-imported"),
    pytest.param("c12", "summarized", {"gap"}, "", "2021-05-31", id="gap-unpaired-trade-summarized"),
    pytest.param("c13", "summarized", {"window"}, "", "2021-05-31", id="border-orphan-segment-2"),
    pytest.param("b1", "pair", {"window"}, "", "", id="segment-2-pair"),
    pytest.param("c15", "standalone", {"window"}, "", "", id="segment-2-standalone"),
    pytest.param("c17", "deferred", {"window", "after"}, "", "", id="beyond-end-of-last-segment-deferred"),
    pytest.param("c18", "deferred", {"after"}, "", "", id="after-unpaired-trade-deferred"),
    pytest.param("c19", "standalone", {"after"}, "", "", id="after-standalone"),
]


class TestZones:
    """Every cell of the zones table, on one set: two segments, a proven gap, cash rows before, between and after.

    Segment 1: trades 2021-03-01…03-31, C1 = 02-28, border 03-01…03-05, window to 04-07.
    Segment 2: trades 06-01…06-30, C2 = 05-31, border 06-01…06-07, window to 07-07.
    """

    @pytest.mark.parametrize(("key", "kind", "zones", "reason", "checkpoint"), ZONE_CELLS)
    def test_cell(self, key: str, kind: str, zones: Set[str], reason: str, checkpoint: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table = _scenario("zones", tmp_path_factory)

        row = builder.row(_records(table), key)

        assert row["lf_row_kind"] == kind
        assert row["lf_zone"] in zones
        assert (row["lf_reason"], row["lf_checkpoint"]) == (reason, checkpoint)

    @pytest.mark.parametrize(("custody", "cash", "match_key"), [("a1", "c6", "2021-03-03|-208.00|EUR"), ("a3", "c9", "2021-04-02|92.00|EUR"), ("b1", "c14", "2021-06-03|-98.00|EUR"), ("b2", "c16", "2021-07-02|-38.00|EUR")])
    def test_pairs(self, custody: str, cash: str, match_key: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table = _scenario("zones", tmp_path_factory)
        records = _records(table)

        row = builder.row(records, custody)

        assert row is builder.row(records, cash)
        assert (row["lf_row_kind"], row["lf_zone"], row["lf_match_key"]) == ("pair", "window", match_key)

    def test_every_member_row_lands_in_exactly_one_data_row(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table = _scenario("zones", tmp_path_factory)

        _assert_each_source_row_once(builder.row_counts, _records(table))

    def test_truth_points(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """C1: 4746.50 at the end of 02-28 plus the border orphan (−120); C2: 4557.00 at the end of 05-31 plus its orphan (−150)."""
        _builder, table, _path, output = _scenario_parse("zones", tmp_path_factory)
        records = _records(table)

        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"], row["lf_zone"]) for row in _truth(records, "truth_cash")] == [("2021-02-28", "4626.50", "EUR", "before"), ("2021-05-31", "4407.00", "EUR", "gap")]
        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"]) for row in _truth(records, "verification")] == [("2021-06-30", "4305.50", "EUR")]
        assert _truth(records, "truth_position") == []
        opening, gap = _checkpoint(output, "2021-02-28"), _checkpoint(output, "2021-05-31")
        assert (opening.kind, _money(opening.cash), opening.absorbed.count, _money(opening.absorbed.opening_cash)) == ("opening", {"EUR": Decimal("4626.50")}, 5, {"EUR": Decimal("0")})
        assert (gap.kind, _money(gap.cash), gap.absorbed.count, _money(gap.absorbed.cash), gap.absorbed.opening_cash) == ("gap", {"EUR": Decimal("4407.00")}, 3, {"EUR": Decimal("-300")}, [])
        assert [(item.as_of, _money(item.cash)) for item in output.verifications] == [(_d("2021-06-30"), {"EUR": Decimal("4305.50")})]


# =============================================================================
# 9. TRUTH POINTS: checkpoints, verification, balance chain
# =============================================================================


class TestTruthPoints:
    """B0 "Punti di verità": where the checkpoints are, their cash, what they absorb, the final verification, the chain."""

    def test_a_late_cash_statement_moves_the_opening(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """A13: C1 is the eve of the first cash row (03-10), its cash the balance before that row; a trade settled before is outside the coverage."""
        builder, table, _path, output = _scenario_parse("late", tmp_path_factory)
        records = _records(table)

        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"]) for row in _truth(records, "truth_cash")] == [("2021-03-09", "1000.00", "EUR")]
        opening = _checkpoint(output, "2021-03-09")
        assert (opening.kind, _money(opening.cash), opening.absorbed.count, opening.absorbed.rows) == ("opening", {"EUR": Decimal("1000.00")}, 0, [])
        assert _money(opening.absorbed.opening_cash) == {"EUR": Decimal("1000.00")}
        early = builder.row(records, "k1")
        assert (early["lf_row_kind"], early["lf_reason"]) == ("excluded", "outside_cash_coverage")
        assert builder.row(records, "t") is builder.row(records, "e1")
        assert builder.row(records, "k2") is builder.row(records, "e2")

    def test_a_trade_settled_in_a_hole_between_cash_statements(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table = _scenario("hole", tmp_path_factory)
        records = _records(table)

        row = builder.row(records, "k")

        assert (row["lf_row_kind"], row["lf_reason"]) == ("excluded", "outside_cash_coverage")
        assert builder.row(records, "a") is builder.row(records, "g2")
        assert builder.row(records, "l") is builder.row(records, "h2")

    def test_a_trade_settled_after_the_last_cash_row_is_not_yet_settled(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table = _scenario("early_end", tmp_path_factory)

        row = builder.row(_records(table), "k")

        assert (row["lf_row_kind"], row["lf_reason"]) == ("excluded", "not_yet_settled")

    def test_the_opening_absorbs_the_rows_before_it(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """``opening_cash`` = the checkpoint cash minus what it absorbs: here the balance the statement starts from."""
        _builder, _table, _path, output = _scenario_parse("early_end", tmp_path_factory)

        opening = _checkpoint(output, "2021-02-28")

        assert (opening.kind, _money(opening.cash), opening.absorbed.count) == ("opening", {"EUR": Decimal("1500.00")}, 1)
        assert [(row.as_of, row.currency, row.amount, row.label) for row in opening.absorbed.rows] == [(_d("2021-02-01"), "EUR", Decimal("500"), "Matti Meikäläinen")]
        assert (_money(opening.absorbed.cash), _money(opening.absorbed.opening_cash)) == ({"EUR": Decimal("500")}, {"EUR": Decimal("1000.00")})

    @pytest.mark.parametrize(
        ("source", "name", "as_of", "amount"),
        [
            pytest.param("sample", "main", "2020-06-26", "1994.46", id="segment-ends-first"),
            pytest.param("scenario", "early_end", "2021-03-25", "1288.50", id="cash-ends-first"),
            pytest.param("scenario", "late", "2021-03-31", "1015.00", id="late-cash-statement"),
        ],
    )
    def test_the_verification_is_at_the_earlier_end(self, source: str, name: str, as_of: str, amount: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """V = min(end of the last segment, last cash row), with the cash at the end of that day."""
        table, _path, output = _combined(source, name, tmp_path_factory)

        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"]) for row in _truth(_records(table), "verification")] == [(as_of, amount, "EUR")]
        assert [(item.as_of, _money(item.cash)) for item in output.verifications] == [(_d(as_of), {"EUR": Decimal(amount)})]

    def test_a_gap_checkpoint_needs_a_proven_gap(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _merged_builder, merged = _scenario("merged", tmp_path_factory)
        _zones_builder, zones = _scenario("zones", tmp_path_factory)

        assert [row["lf_checkpoint"] for row in _truth(_records(merged), "truth_cash")] == ["2021-02-28"]
        assert [row["lf_checkpoint"] for row in _truth(_records(zones), "truth_cash")] == ["2021-02-28", "2021-05-31"]

    def test_the_rows_of_an_unproven_gap_stay_in_the_window(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        builder, table, _path, output = _scenario_parse("merged", tmp_path_factory)
        records = _records(table)

        for key in ("deposit", "fee"):
            row = builder.row(records, key)
            assert (row["lf_row_kind"], row["lf_zone"], row["lf_checkpoint"]) == ("standalone", "window", ""), key
        assert [(item.as_of, item.kind, _money(item.cash)) for item in output.checkpoints] == [(_d("2021-02-28"), "opening", {"EUR": Decimal("5000.00")})]

    def test_a_broken_balance_chain_turns_the_checkpoint_cash_into_a_verification(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The truth_cash row stays, with no value; the checkpoint has no cash; a verification at its date carries the stated balance."""
        _builder, table, _path, output = _scenario_parse("chain", tmp_path_factory)
        records = _records(table)

        assert [(row["lf_checkpoint"], row["lf_value"]) for row in _truth(records, "truth_cash")] == [("2021-02-28", "")]
        assert ("2021-02-28", "1500.00") in [(row["lf_checkpoint"], row["lf_value"]) for row in _truth(records, "verification")]
        assert _checkpoint(output, "2021-02-28").cash == []
        assert (_d("2021-02-28"), {"EUR": Decimal("1500.00")}) in [(item.as_of, _money(item.cash)) for item in output.verifications]
        notices = [notice for notice in output.warnings if notice.code == "balance_chain_broken"]
        assert notices and all(notice.severity == "warning" for notice in notices)


# =============================================================================
# 10. PROOFS OF POSITION (E1–E4)
# =============================================================================


class TestProofs:
    """B0: one position per title and checkpoint, from the proofs of the segment, carried back to C_i."""

    @pytest.mark.parametrize("name", ["recent_trade", "recent_cash"])
    def test_e1_is_discarded_by_a_trade_in_the_30_days_before(self, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """A purchase of the title (custody row, or an unpaired ``Osto`` cash row) in the 30 days before the ``Tuotto``."""
        _builder, table, _path, output = _scenario_parse(name, tmp_path_factory)

        assert [row["lf_proof"] for row in _positions_by_title(_records(table), "Aurora Metsä Oyj")] == ["discarded:E1:recent_trades"]
        assert "Aurora Metsä Oyj" not in _positions(output, _checkpoint(output, "2021-02-28"))
        assert "proof_discarded" in _codes(output.warnings)

    def test_e1_is_discarded_when_the_30_days_fall_before_the_cash_statement(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _builder, table, _path, output = _scenario_parse("late", tmp_path_factory)

        assert [row["lf_proof"] for row in _positions_by_title(_records(table), "Aurora Metsä Oyj")] == ["discarded:E1:window_not_covered"]
        assert "Aurora Metsä Oyj" not in _positions(output, _checkpoint(output, "2021-03-09"))
        assert "proof_discarded" in _codes(output.warnings)

    def test_e2_is_discarded_by_another_movement_the_same_day(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Without an exact proof E3 applies: a sale of 10 and the old line of 100 prove at least 110."""
        _builder, table, _path, output = _scenario_parse("same_day", tmp_path_factory)
        proofs = [row["lf_proof"] for row in _positions_by_title(_records(table), "Lumi Holding Oyj")]

        assert "discarded:E2:same_day_movement" in proofs
        assert "exact:E2" not in proofs
        assert _positions(output, _checkpoint(output, "2021-02-28"))["Lumi Holding Oyj"] == (Decimal("110"), "at_least")
        assert "proof_discarded" in _codes(output.warnings)

    def test_e4_an_excluded_row_before_the_proof_discards_it(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """A cancelled purchase of the title between C1 and the ``Tuotto``: the ``Tuotto`` proves nothing."""
        _builder, table, _path, output = _scenario_parse("excluded_proof", tmp_path_factory)
        proofs = [row["lf_proof"] for row in _positions_by_title(_records(table), "Aurora Metsä Oyj")]

        assert proofs and all(re.fullmatch(r"discarded:E[1-4]:excluded_rows", proof) for proof in proofs), proofs
        assert "Aurora Metsä Oyj" not in _positions(output, _checkpoint(output, "2021-02-28"))
        assert "proof_discarded" in _codes(output.warnings)

    def test_e4_a_superseded_row_does_not_discard(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The row rule M excluded as ``superseded`` is not an excluded row for E4: the ``Tuotto`` proves exactly 60."""
        builder, table, _path, output = _scenario_parse("superseded_proof", tmp_path_factory)
        records = _records(table)

        assert (builder.row(records, "x")["lf_row_kind"], builder.row(records, "x")["lf_reason"]) == ("excluded", "superseded")
        assert [(row["lf_proof"], row["lf_value"]) for row in _positions_by_title(records, "Aurora Metsä Oyj")] == [("exact:E1", "60")]
        assert _positions(output, _checkpoint(output, "2021-02-28"))["Aurora Metsä Oyj"] == (Decimal("60"), "exact")

    def test_exact_proofs_that_disagree_are_all_discarded(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The ``Tuotto`` says 150, the demerger 120, nothing moved in between: both go (``conflict``); E3 still proves 120."""
        _builder, table, _path, output = _scenario_parse("conflict", tmp_path_factory)
        proofs = [row["lf_proof"] for row in _positions_by_title(_records(table), "Lumi Holding Oyj")]

        assert {"discarded:E1:conflict", "discarded:E2:conflict"} <= set(proofs)
        assert not [proof for proof in proofs if proof.startswith("exact:")]
        assert _positions(output, _checkpoint(output, "2021-02-28"))["Lumi Holding Oyj"] == (Decimal("120"), "at_least")
        assert "proof_discarded" in _codes(output.warnings)


# =============================================================================
# 11. THE COMBINED FORMAT
# =============================================================================


class TestCombinedFormat:
    """B0 "Combinato" and design §3.4.4: headers, value formats, order, verbatim cells, purity."""

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_headers(self, set_name: str) -> None:
        table = _sample_table(set_name)

        assert isinstance(table, BRIMCombinedTable)
        assert table.headers == COMBINED_HEADERS
        assert set(COMBINED_REQUIRED_HEADERS) <= set(table.headers)

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_value_formats(self, set_name: str) -> None:
        records = _records(_sample_table(set_name))
        first_checkpoint = min(row["lf_checkpoint"] for row in _truth(records, "truth_cash"))

        for number, record in enumerate(records, start=2):
            where = f"line {number} ({record['lf_row_kind']} {record['lf_source']})"
            kind = record["lf_row_kind"]
            assert kind in DATA_KINDS | TRUTH_KINDS, where
            assert kind == "verification" or record["lf_zone"] in ZONES, where
            assert record["lf_reason"] in (REASONS if kind == "excluded" else {""}), where
            assert record["lf_checkpoint"] == "" or ISO_DATE_RE.match(record["lf_checkpoint"]), where
            _sources(record)
            if kind in DATA_KINDS:
                assert (record["lf_value"], record["lf_proof"]) == ("", ""), where
                assert record["lf_match_key"] == "" or MATCH_KEY_RE.match(record["lf_match_key"]), where
            if kind == "pair":
                assert record["lf_match_key"] == f"{_from_fi_date(record['cash:Pvm']).isoformat()}|{_number(record['cash:Määrä']):.2f}|EUR", where
            if kind in TRUTH_KINDS:
                assert ISO_DATE_RE.match(record["lf_checkpoint"]), where
            if kind in ("truth_cash", "verification"):
                assert record["lf_currency"] == "EUR" and CASH_VALUE_RE.match(record["lf_value"]), where
            if kind == "truth_position":
                assert PROOF_RE.match(record["lf_proof"]), where
                assert record["lf_proof"].startswith("discarded:") or QUANTITY_VALUE_RE.match(record["lf_value"]), where
            if kind in ("truth_cash", "truth_position"):
                assert record["lf_zone"] == ("before" if record["lf_checkpoint"] == first_checkpoint else "gap"), where

    def test_source_tokens(self) -> None:
        """``custody:12``, ``cash:40``, ``custody:12 + cash:40``; ``custody#2:12`` for the files of a role that has several."""
        tokens = {name: {token.split(":")[0] for record in _records(_sample_table(name)) for token in record["lf_source"].split(" + ")} for name in sorted(SAMPLE_SETS)}

        assert tokens == {"gap": {"custody#1", "custody#2", "cash"}, "main": {"custody", "cash"}}
        assert _data_row(_records(_sample_table("main")), "custody", 1, 2)["lf_source"] == "custody:2 + cash:18"
        assert _data_row(_records(_sample_table("gap")), "custody", 2, 4)["lf_source"] == "custody#2:4 + cash:4"

    def test_member_cells_are_copied_verbatim(self) -> None:
        """XLSX text as is, XLSX numbers canonical ("50", "41.2", "-1181.04"), dates dd.mm.yyyy, CSV cells as is."""
        records = _records(_sample_table("main"))

        aurora = _data_row(records, "custody", 1, 2)
        assert {header: aurora[header] for header in COMBINED_HEADERS[len(LF_HEADERS) :]} == {
            "custody:Kauppapäivä": "10.02.2020",
            "custody:Arvopäivä": "12.02.2020",
            "custody:Sijoituskohde": "Aurora Metsä Oyj",
            "custody:Määrä": "50",
            "custody:Kurssi": "14.2",
            "custody:Palkkio<br/>sis. Alv": "0",
            "custody:Summa": "-718",
            "custody:H": "EUR",
            "custody:Tila": "Toteutettu",
            "custody:Toimeksiantotyyppi": "Rajakurssi",
            "cash:Pvm": "12.02.2020",
            "cash:Saaja/Maksaja": "Osto Aurora Metsä Oyj",
            "cash:Määrä": "-718",
            "cash:Saldo": "800,46",
            "cash:Tila": "Toteutunut",
            "cash:Tarkastus": "Ei",
        }
        kaamos = _data_row(records, "custody", 1, 5)
        assert (kaamos["custody:Määrä"], kaamos["custody:Kurssi"], kaamos["custody:Summa"], kaamos["cash:Määrä"], kaamos["cash:Saldo"]) == ("300", "41.2", "-1181.04", "-1181,04", "1518,46")
        tuotto = _data_row(records, "custody", 1, 3)
        assert (tuotto["custody:Kurssi"], tuotto["custody:Summa"], tuotto["cash:Määrä"], tuotto["cash:Saaja/Maksaja"]) == ("0", "22.5", "22,5", "Aurora Metsä 4820193756")
        vanha = _data_row(records, "custody", 1, 7)
        assert (vanha["custody:Määrä"], vanha["custody:Summa"], vanha["custody:H"], vanha["custody:Toimeksiantotyyppi"], vanha["cash:Pvm"]) == ("-100", "", "", "Jakautuminen, vanha", "")
        assert _data_row(records, "custody", 1, 16)["custody:Kurssi"] == "21"
        fee = _data_row(records, "cash", 1, 3)
        assert [fee[header] for header in COMBINED_HEADERS if header.startswith("custody:")] == [""] * 10
        assert (fee["cash:Pvm"], fee["cash:Saaja/Maksaja"], fee["cash:Määrä"], fee["cash:Saldo"]) == ("30.06.2020", "Palvelumaksut 06/2020", "-3,5", "1772,96")

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_rows_are_in_order(self, set_name: str) -> None:
        """By value date (truth rows at their date, after that day's data rows), then custody before cash, then file and line."""
        records = _records(_sample_table(set_name))
        keys = [_sort_key(record) for record in records]

        inversions = [f"line {index + 2} ({records[index]['lf_source']}) after line {index + 1} ({records[index - 1]['lf_source']})" for index in range(1, len(keys)) if keys[index] < keys[index - 1]]

        assert inversions == []

    @pytest.mark.parametrize(("set_name", "row_counts"), [("main", {("custody", 1): 15, ("cash", 1): 32}), ("gap", {("custody", 1): 3, ("custody", 2): 4, ("cash", 1): 14})])
    def test_every_member_row_lands_in_exactly_one_data_row(self, set_name: str, row_counts: Dict[Tuple[str, int], int]) -> None:
        _assert_each_source_row_once(row_counts, _records(_sample_table(set_name)))

    def test_combine_is_pure(self) -> None:
        """Same members, same table (D-S26); the members and the argument are left as they were."""
        plugin = _danske()
        members = _members(GAP_SET)
        before = {path: path.read_bytes() for paths in members.values() for path in paths}

        first, second = plugin.combine(members), plugin.combine(_members(GAP_SET))

        assert (first.headers, first.rows, first.summary) == (second.headers, second.rows, second.summary)
        assert {path: path.read_bytes() for path in before} == before
        assert members == _members(GAP_SET)

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_the_summary_can_be_stored_in_the_sidecar(self, set_name: str) -> None:
        """The core writes the summary into the combined file's JSON sidecar: plain JSON values only."""
        summary = _sample_table(set_name).summary

        assert summary
        try:
            json.dumps(summary)
        except TypeError as error:
            pytest.fail(f"the combine summary is not plain JSON: {error}")

    def test_written_by_the_core_it_reads_back_unchanged(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """No cell breaks the one-row-one-line layout the evidence line numbers rely on."""
        table, path, _output = _sample_parse("gap", tmp_path_factory)

        assert path.read_bytes().startswith(b"\xef\xbb\xbf")
        with open(path, encoding="utf-8-sig", newline="") as handle:
            assert list(csv.reader(handle, delimiter=";")) == [table.headers, *table.rows]
        assert _line_count(path) == len(table.rows) + 1


class TestPrivacy:
    """B0: ``Säilytystili`` is an account number; it serves the in-memory check and never leaves the members."""

    @pytest.mark.parametrize("set_name", sorted(SAMPLE_SETS))
    def test_samples(self, set_name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        table, path, output = _sample_parse(set_name, tmp_path_factory)

        assert [header for header in table.headers if "Säilytystili" in header] == []
        assert ACCOUNT not in path.read_text(encoding="utf-8-sig")
        assert ACCOUNT not in json.dumps(table.summary, default=str)
        assert ACCOUNT not in output.model_dump_json()

    def test_a_synthetic_set(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _builder, table, path, output = _scenario_parse("zones", tmp_path_factory)

        assert ACCOUNT not in path.read_text(encoding="utf-8-sig") + json.dumps(table.summary, default=str) + output.model_dump_json()


# =============================================================================
# 12. THE PARSE OF THE COMBINED FILE — contract, on the samples and on synthetic sets
# =============================================================================

PARSED_SETS = [
    pytest.param("sample", "main", id="main"),
    pytest.param("sample", "gap", id="gap"),
    pytest.param("scenario", "zones", id="zones"),
    pytest.param("scenario", "classes", id="classes"),
    pytest.param("scenario", "same_day", id="same_day"),
    pytest.param("scenario", "conflict", id="conflict"),
]


class TestParseContract:
    """B0 "Parse del combinato": what the parse owes for every row, todo, notice and truth point of the file."""

    @pytest.mark.parametrize("name", sorted(SAMPLE_ROLES))
    def test_a_member_alone_is_refused(self, name: str) -> None:
        """A second guard behind the core's: a member is imported through its combined file only."""
        plugin = _danske()

        with pytest.raises(BRIMSetRequiredError):
            plugin.parse(SAMPLE_DIR / name, broker_id=BROKER_ID)

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_no_validation_issue(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _combined(source, name, tmp_path_factory)

        assert output.validation_issues == []
        assert all(tx.broker_id == BROKER_ID for tx in output.transactions)

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_transactions_follow_the_combined_rows(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """One per pair and standalone row, in file order: type, value date, quantity, cash, title, description and tags."""
        table, _path, output = _combined(source, name, tmp_path_factory)

        assert _transactions(output) == _owed_transactions(_records(table))
        assert all(tx.cash is None or tx.cash.code == "EUR" for tx in output.transactions)

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_parse_is_idempotent(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, path, output = _combined(source, name, tmp_path_factory)

        assert _danske().parse(path, broker_id=BROKER_ID).model_dump(mode="json") == output.model_dump(mode="json")

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_one_fake_asset_per_title(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The custody name is the extracted name; titles that only a truth position names get an id too."""
        table, _path, output = _combined(source, name, tmp_path_factory)
        records = _records(table)
        names = _names(output)

        assert all(is_fake_asset_id(asset_id) for asset_id in names)
        assert all(title and title.strip() for title in names.values())
        assert len(set(names.values())) == len(names), "one fake id per title"
        used = {tx.asset_id for tx in output.transactions if tx.asset_id is not None} | {position.asset_id for checkpoint in output.checkpoints for position in checkpoint.positions}
        assert used <= set(names)
        titles = {record["custody:Sijoituskohde"] for record in records if record["lf_row_kind"] in ("pair", "standalone") and record["custody:Sijoituskohde"]}
        titles |= {record["custody:Sijoituskohde"] for record in _truth(records, "truth_position") if not record["lf_proof"].startswith("discarded:")}
        assert titles <= set(names.values())

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_todos_per_transaction(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Every BUY and SELL: charges included (warning on ``cash``). A demerger: a blocking cost on the new lines, a warning on the old one."""
        _table, _path, output = _combined(source, name, tmp_path_factory)
        by_transaction: Dict[int, List[Tuple[str, str, str]]] = defaultdict(list)
        for todo in output.field_todos:
            by_transaction[todo.tx_index].append((todo.reason_code, todo.field, todo.severity))

        assert set(by_transaction) <= set(range(len(output.transactions)))
        for index, tx in enumerate(output.transactions):
            if tx.type.value in ("BUY", "SELL"):
                expected = [(CHARGES, "cash", "warning")]
            elif tx.type.value == "ADJUSTMENT":
                expected = [(DEMERGER, "cost_basis_override", "blocker")] if tx.quantity > 0 else [(DEMERGER_OLD, "quantity", "warning")]
            else:
                expected = []
            assert sorted(by_transaction.get(index, [])) == expected, f"transaction {index}: {tx.type.value} {tx.date} {tx.description}"

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_charges_todo_context(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The split channel: row (a line of the combined file), cash, currency, no nominal, the hint, suggestions — and the charges when plausible in EUR."""
        table, _path, output = _combined(source, name, tmp_path_factory)
        records = _records(table)
        names = _names(output)
        todos = [todo for todo in output.field_todos if todo.reason_code == CHARGES]
        assert todos or not any(tx.type.value in ("BUY", "SELL") for tx in output.transactions)

        for todo in todos:
            tx = output.transactions[todo.tx_index]
            context = todo.context or {}
            where = f"{tx.type.value} {tx.date} {tx.description}"
            assert set(CHARGES_CONTEXT) <= set(context), where
            assert isinstance(context["row"], int) and 2 <= context["row"] <= len(records) + 1, where
            record = records[context["row"] - 2]
            assert (record["lf_row_kind"], _from_fi_date(record["cash:Pvm"]), record["custody:Sijoituskohde"]) == ("pair", tx.date, names[tx.asset_id]), where
            assert context["cash"] == f"{abs(tx.cash.amount):.2f}", f"{where}: cash is the absolute amount with two decimals"
            assert (context["currency"], context["compare_nominal"], context["split_hint"]) == ("EUR", False, "trade_charges"), where
            suggestions = context["split_suggestions"]
            assert isinstance(suggestions, list) and suggestions and all(isinstance(item, str) and item.strip() for item in suggestions), where
            gross = abs(_number(record["custody:Määrä"]) * _number(record["custody:Kurssi"]))
            difference = abs(tx.cash.amount) - gross if tx.quantity > 0 else gross - abs(tx.cash.amount)
            if Decimal(0) <= difference <= max(Decimal(15), gross / 100):
                assert context.get("charges") == f"{difference:.2f}", where
            else:
                assert "charges" not in context, f"{where}: the price is in another currency, there is no amount to suggest"

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_every_todo_shows_its_source_rows(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, path, output = _combined(source, name, tmp_path_factory)
        lines = _line_count(path)

        for todo in output.field_todos:
            where = f"{todo.reason_code} on transaction {todo.tx_index}"
            assert todo.message and todo.message.strip(), where
            assert todo.evidence, where
            for evidence in todo.evidence:
                assert evidence.title.strip() and evidence.comment and evidence.comment.strip(), where
                assert evidence.row_numbers and len(evidence.row_numbers) == len(evidence.rows), where
                assert all(1 <= number <= lines for number in evidence.row_numbers), where
                assert all(len(row) == len(evidence.headers) for row in evidence.rows), where

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_notices(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Stable codes with their severity; one ``excluded_<reason>`` per reason, showing exactly those rows; the others when they apply."""
        table, path, output = _combined(source, name, tmp_path_factory)
        records = _records(table)
        lines = _line_count(path)
        excluded: Dict[str, Set[int]] = defaultdict(set)
        for number, record in enumerate(records, start=2):
            if record["lf_row_kind"] == "excluded":
                excluded[record["lf_reason"]].add(number)
        by_code: Dict[str, List[Any]] = defaultdict(list)
        for notice in output.warnings:
            by_code[notice.code].append(notice)

        for code, notices in by_code.items():
            severity = "warning" if code.startswith("excluded_") else PARSE_NOTICES.get(code)
            assert severity, f"unknown parse notice code {code!r}"
            for notice in notices:
                assert notice.severity == severity, code
                assert notice.message and notice.message.strip(), code
                for evidence in notice.evidence:
                    assert evidence.title.strip() and evidence.comment and evidence.comment.strip(), code
                    assert all(1 <= number <= lines for number in evidence.row_numbers), code
        assert sorted(code for code in by_code if code.startswith("excluded_")) == sorted(f"excluded_{reason}" for reason in excluded)
        for reason, numbers in excluded.items():
            notices = by_code[f"excluded_{reason}"]
            assert len(notices) == 1, reason
            assert {number for evidence in notices[0].evidence for number in evidence.row_numbers} == numbers, reason
        assert ("deferred_rows" in by_code) == any(record["lf_row_kind"] == "deferred" for record in records)
        assert ("proof_discarded" in by_code) == any(record["lf_proof"].startswith("discarded:") for record in _truth(records, "truth_position"))
        assert ("deposit_assumed" in by_code) == any(tx.type.value == "DEPOSIT" for tx in output.transactions)

    @pytest.mark.parametrize(("source", "name"), PARSED_SETS)
    def test_checkpoints_and_verifications_are_the_truth_rows(self, source: str, name: str, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The earliest checkpoint is the opening; positions carry the transactions' fake ids; the absorbed rows are the data rows of that date."""
        table, _path, output = _combined(source, name, tmp_path_factory)
        records = _records(table)
        names = _names(output)
        truth_cash = _truth(records, "truth_cash")
        dates = sorted(row["lf_checkpoint"] for row in truth_cash)

        assert sorted(checkpoint.as_of.isoformat() for checkpoint in output.checkpoints) == dates
        for checkpoint in output.checkpoints:
            day = checkpoint.as_of.isoformat()
            [row] = [item for item in truth_cash if item["lf_checkpoint"] == day]
            assert checkpoint.kind == ("opening" if day == dates[0] else "gap"), day
            assert _money(checkpoint.cash) == ({row["lf_currency"]: Decimal(row["lf_value"])} if row["lf_value"] else {}), day
            proven = sorted((item["custody:Sijoituskohde"], Decimal(item["lf_value"]), item["lf_proof"].split(":")[0]) for item in _truth(records, "truth_position", day) if not item["lf_proof"].startswith("discarded:"))
            assert sorted((names.get(position.asset_id), position.quantity, position.exactness) for position in checkpoint.positions) == proven, day
            absorbed = sorted((_from_fi_date(item["cash:Pvm"]), "EUR", _number(item["cash:Määrä"]), item["cash:Saaja/Maksaja"]) for item in records if item["lf_row_kind"] in DATA_KINDS and item["lf_checkpoint"] == day and item["cash:Määrä"])
            assert sorted((item.as_of, item.currency, item.amount, item.label) for item in checkpoint.absorbed.rows) == absorbed, day
            assert checkpoint.absorbed.count == len(absorbed), day
            total = sum((item[2] for item in absorbed), Decimal(0))
            assert _nonzero(_money(checkpoint.absorbed.cash)) == _nonzero({"EUR": total}), day
            if checkpoint.kind == "gap":
                assert checkpoint.absorbed.opening_cash == [], day
            elif checkpoint.cash:
                assert _money(checkpoint.absorbed.opening_cash) == {"EUR": _money(checkpoint.cash)["EUR"] - total}, day
        stated = sorted((row["lf_checkpoint"], row["lf_currency"], Decimal(row["lf_value"])) for row in _truth(records, "verification"))
        assert sorted((item.as_of.isoformat(), cash.currency, cash.amount) for item in output.verifications for cash in item.cash) == stated


# =============================================================================
# 13. THE MAIN SAMPLE SET
# =============================================================================

# (type, value date, quantity, cash, title, description), in the order of the combined file.
MAIN_TRANSACTIONS = [
    ("DEPOSIT", "2019-01-07", "0", "6000", None, "Matti Meikäläinen"),
    ("FEE", "2019-03-29", "0", "-3.5", None, "Palvelumaksut 03/2019"),
    ("FEE", "2019-06-28", "0", "-3.5", None, "Palvelumaksut 06/2019"),
    ("FEE", "2019-09-30", "0", "-3.5", None, "Palvelumaksut 09/2019"),
    ("TAX", "2019-11-04", "0", "-21.3", None, "Vero osakesäästötililtä"),
    ("WITHDRAWAL", "2019-11-04", "0", "-300", None, "Nosto osakesäästötililtä"),
    ("FEE", "2019-12-31", "0", "-3.5", None, "Palvelumaksut 12/2019"),
    ("BUY", "2020-02-05", "300", "-1181.04", "Kaamos Robotics AB", "Osto Kaamos Robotics AB (Rajakurssi)"),
    ("BUY", "2020-02-12", "50", "-718", "Aurora Metsä Oyj", "Osto Aurora Metsä Oyj (Rajakurssi)"),
    ("BUY", "2020-03-04", "40", "-402", "Lumi Holding Oyj", "Osto Lumi Holding Oyj (Pikakauppa)"),
    ("SELL", "2020-03-12", "-120", "724", "Pohjola Tekniikka Oyj", "Myynti Pohjola Tekniikka (Päivän kurssi)"),
    ("DEPOSIT", "2020-03-20", "0", "1000", None, "Matti Meikäläinen"),
    ("BUY", "2020-03-27", "25", "-318", "Saimaa Energia Oyj", "Osto Saimaa Energia Oyj (Rajakurssi)"),
    ("BUY", "2020-03-27", "25", "-318", "Saimaa Energia Oyj", "Osto Saimaa Energia Oyj (Rajakurssi)"),
    ("FEE", "2020-03-31", "0", "-3.5", None, "Palvelumaksut 03/2020"),
    ("BUY", "2020-04-03", "20", "-116", "Pohjola Tekniikka Oyj", "Osto Pohjola Tekniikka O (Rajakurssi)"),
    ("DIVIDEND", "2020-04-15", "0", "22.5", "Aurora Metsä Oyj", "Aurora Metsä 4820193756 (Tuotto)"),
    ("SELL", "2020-05-07", "-80", "532", "Pohjola Tekniikka Oyj", "Myynti Pohjola Tekniikka (Päivän kurssi)"),
    ("SELL", "2020-05-22", "-30", "466", "Aurora Metsä Oyj", "Myynti Aurora Metsä Oyj (Rajakurssi)"),
    ("ADJUSTMENT", "2020-06-01", "-100", None, "Lumi Holding Oyj", "Jakautuminen, vanha: Lumi Holding Oyj"),
    ("ADJUSTMENT", "2020-06-01", "100", None, "Lumi Kiinteistöt Oyj", "Jakautuminen, uusi: Lumi Kiinteistöt Oyj"),
    ("ADJUSTMENT", "2020-06-01", "100", None, "Lumi Palvelut Oyj", "Jakautuminen, uusi: Lumi Palvelut Oyj"),
    ("TAX", "2020-06-10", "0", "-45", None, "Vero osakesäästötililtä"),
    ("WITHDRAWAL", "2020-06-10", "0", "-500", None, "Nosto osakesäästötililtä"),
    ("SELL", "2020-06-17", "-50", "152", "Lumi Palvelut Oyj", "Myynti Lumi Palvelut Oyj (Päivän kurssi)"),
    ("BUY", "2020-06-30", "10", "-218", "Tunturi Pankki Oyj", "Osto Tunturi Pankki Oyj (Rajakurssi)"),
    ("FEE", "2020-06-30", "0", "-3.5", None, "Palvelumaksut 06/2020"),
]
# custody line -> cash line of every pair but the two identical partial fills (custody 14 and 15, cash 13 and 14).
MAIN_PAIRS = {2: 18, 3: 10, 4: 8, 5: 19, 6: 17, 10: 5, 11: 16, 12: 11, 13: 9, 16: 4}
# cash line, row kind, accepted zones, checkpoint (C1 = 2020-02-02 absorbs every cash row up to it, and the border orphan).
MAIN_CASH_ROWS = [
    pytest.param(2, "deferred", {"window", "after"}, "", id="purchase-settled-after-the-custody-export"),
    pytest.param(3, "standalone", {"window"}, "", id="fee-06-2020"),
    pytest.param(6, "standalone", {"window"}, "", id="tax-2020"),
    pytest.param(7, "standalone", {"window"}, "", id="withdrawal-2020"),
    pytest.param(12, "standalone", {"window"}, "", id="fee-03-2020"),
    pytest.param(15, "standalone", {"window"}, "", id="deposit-2020"),
    pytest.param(20, "summarized", {"window"}, "2020-02-02", id="border-orphan"),
    pytest.param(21, "standalone", {"before"}, "2020-02-02", id="fee-12-2019"),
    pytest.param(22, "standalone", {"before"}, "2020-02-02", id="tax-2019"),
    pytest.param(23, "standalone", {"before"}, "2020-02-02", id="withdrawal-2019"),
    pytest.param(24, "summarized", {"before"}, "2020-02-02", id="sale-2019"),
    pytest.param(25, "standalone", {"before"}, "2020-02-02", id="fee-09-2019"),
    pytest.param(26, "standalone", {"before"}, "2020-02-02", id="fee-06-2019"),
    pytest.param(27, "summarized", {"before"}, "2020-02-02", id="income-2019"),
    pytest.param(28, "standalone", {"before"}, "2020-02-02", id="fee-03-2019"),
    pytest.param(29, "summarized", {"before"}, "2020-02-02", id="purchase-pohjola-2019"),
    pytest.param(30, "summarized", {"before"}, "2020-02-02", id="purchase-revontuli-2019"),
    pytest.param(31, "summarized", {"before"}, "2020-02-02", id="purchase-lumi-2019"),
    pytest.param(32, "summarized", {"before"}, "2020-02-02", id="purchase-aurora-2019"),
    pytest.param(33, "standalone", {"before"}, "2020-02-02", id="first-deposit-2019"),
]


def _pair_lines(records: List[Dict[str, str]]) -> Dict[Tuple[int, int], int]:
    """``{(custody file, custody line): cash line}`` of every pair; each pair joins one row of each role."""
    pairs = {}
    for record in records:
        if record["lf_row_kind"] != "pair":
            continue
        custody = [(number, line) for role, number, line in _sources(record) if role == "custody"]
        cash = [line for role, _number, line in _sources(record) if role == "cash"]
        assert len(custody) == len(cash) == 1, record["lf_source"]
        assert (record["lf_zone"], record["lf_reason"], record["lf_checkpoint"]) == ("window", "", ""), record["lf_source"]
        pairs[custody[0]] = cash[0]
    return pairs


def _expected_todos(transactions: Sequence[Tuple[str, ...]]) -> List[Tuple[Any, ...]]:
    """(transaction, reason, field, severity, cash, charges) of the sample todos.

    A split todo states the trade's cash as its absolute amount with two decimals ("718.00") and, for a
    title priced in EUR, charges of 8.00; the title priced in another currency gets no charges.
    """
    todos = []
    for index, (tx_type, _day, quantity, cash, title, _description) in enumerate(transactions):
        if tx_type in ("BUY", "SELL"):
            todos.append((index, CHARGES, "cash", "warning", f"{abs(Decimal(cash)):.2f}", None if title == "Kaamos Robotics AB" else "8.00"))
        elif tx_type == "ADJUSTMENT":
            todos.append((index, DEMERGER, "cost_basis_override", "blocker", None, None) if Decimal(quantity) > 0 else (index, DEMERGER_OLD, "quantity", "warning", None, None))
    return todos


def _sample_todos(output: BRIMParseOutput) -> List[Tuple[Any, ...]]:
    """(transaction, reason, field, severity, cash, charges) of the parse, cash and charges read on the split todos only."""
    todos = []
    for todo in output.field_todos:
        context = (todo.context or {}) if todo.reason_code == CHARGES else {}
        todos.append((todo.tx_index, todo.reason_code, todo.field, todo.severity, context.get("cash"), context.get("charges")))
    return sorted(todos, key=lambda item: (item[0], item[1]))


class TestMainSampleFacts:
    """The main set: a custody export 2020-02-03…06-26 and a Latin-1 cash statement 2019-01-07…2020-07-03."""

    def test_row_kinds(self) -> None:
        kinds = Counter(record["lf_row_kind"] for record in _records(_sample_table("main")))

        assert kinds == Counter({"pair": 12, "standalone": 15, "summarized": 7, "deferred": 1, "truth_cash": 1, "truth_position": 3, "verification": 1})

    def test_pairs(self) -> None:
        pairs = _pair_lines(_records(_sample_table("main")))

        assert {line: cash for (_file, line), cash in pairs.items() if line not in (14, 15)} == MAIN_PAIRS
        assert sorted(pairs[(1, line)] for line in (14, 15)) == [13, 14]

    @pytest.mark.parametrize("line", [7, 8, 9], ids=["old-line", "new-line-1", "new-line-2"])
    def test_demerger_lines_are_standalone(self, line: int) -> None:
        row = _data_row(_records(_sample_table("main")), "custody", 1, line)

        assert (row["lf_row_kind"], row["lf_zone"], row["lf_reason"], row["lf_checkpoint"]) == ("standalone", "window", "", "")

    @pytest.mark.parametrize(("line", "kind", "zones", "checkpoint"), MAIN_CASH_ROWS)
    def test_cash_rows(self, line: int, kind: str, zones: Set[str], checkpoint: str) -> None:
        row = _data_row(_records(_sample_table("main")), "cash", 1, line)

        assert row["lf_row_kind"] == kind
        assert row["lf_zone"] in zones
        assert (row["lf_reason"], row["lf_checkpoint"]) == ("", checkpoint)

    def test_truth_rows(self) -> None:
        """C1 = 2020-02-02: 2819.50 at the end of the day, minus the border orphan (−120 settled 2020-02-03)."""
        records = _records(_sample_table("main"))

        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"], row["lf_zone"]) for row in _truth(records, "truth_cash")] == [("2020-02-02", "2699.50", "EUR", "before")]
        assert sorted((row["custody:Sijoituskohde"], row["lf_value"], row["lf_proof"], row["lf_checkpoint"], row["lf_zone"]) for row in _truth(records, "truth_position")) == [
            ("Aurora Metsä Oyj", "100", "exact:E1", "2020-02-02", "before"),
            ("Lumi Holding Oyj", "60", "exact:E2", "2020-02-02", "before"),
            ("Pohjola Tekniikka Oyj", "180", "at_least:E3", "2020-02-02", "before"),
        ]
        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"]) for row in _truth(records, "verification")] == [("2020-06-26", "1994.46", "EUR")]

    def test_transactions(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("main", tmp_path_factory)

        assert _transactions(output) == [_expected_tx(*row) for row in MAIN_TRANSACTIONS]
        assert Counter(tx.type.value for tx in output.transactions) == Counter({"BUY": 7, "SELL": 4, "DIVIDEND": 1, "ADJUSTMENT": 3, "FEE": 6, "TAX": 2, "WITHDRAWAL": 2, "DEPOSIT": 2})

    def test_assets(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("main", tmp_path_factory)

        assert sorted(_names(output).values()) == ["Aurora Metsä Oyj", "Kaamos Robotics AB", "Lumi Holding Oyj", "Lumi Kiinteistöt Oyj", "Lumi Palvelut Oyj", "Pohjola Tekniikka Oyj", "Saimaa Energia Oyj", "Tunturi Pankki Oyj"]

    def test_todos(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """11 trades with charges included (8.00 each, none for the title priced in another currency), two blocking new lines, one old line."""
        _table, _path, output = _sample_parse("main", tmp_path_factory)

        todos = _sample_todos(output)

        assert todos == _expected_todos(MAIN_TRANSACTIONS)
        assert Counter(todo[1] for todo in todos) == Counter({CHARGES: 11, DEMERGER: 2, DEMERGER_OLD: 1})
        assert [todo[4] for todo in todos if todo[0] in (7, 8)] == ["1181.04", "718.00"]

    def test_notices(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("main", tmp_path_factory)

        assert sorted(set(_codes(output.warnings))) == ["deferred_rows", "deposit_assumed"]

    def test_checkpoint_and_verification(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("main", tmp_path_factory)

        assert len(output.checkpoints) == 1
        opening = output.checkpoints[0]
        assert (opening.as_of, opening.kind, _money(opening.cash)) == (_d("2020-02-02"), "opening", {"EUR": Decimal("2699.50")})
        assert _positions(output, opening) == {"Aurora Metsä Oyj": (Decimal("100"), "exact"), "Lumi Holding Oyj": (Decimal("60"), "exact"), "Pohjola Tekniikka Oyj": (Decimal("180"), "at_least")}
        assert all(position.unit_cost is None for position in opening.positions)
        absorbed = opening.absorbed
        assert (absorbed.count, len(absorbed.rows), _money(absorbed.cash), _money(absorbed.opening_cash)) == (14, 14, {"EUR": Decimal("2699.50")}, {"EUR": Decimal("0")})
        assert (_d("2020-02-03"), "EUR", Decimal("-120"), "Osto Pohjola Tekniikka O") in [(row.as_of, row.currency, row.amount, row.label) for row in absorbed.rows]
        assert (min(row.as_of for row in absorbed.rows), max(row.as_of for row in absorbed.rows)) == (_d("2019-01-07"), _d("2020-02-03"))
        assert [(item.as_of, _money(item.cash)) for item in output.verifications] == [(_d("2020-06-26"), {"EUR": Decimal("1994.46")})]


# =============================================================================
# 14. THE GAP SAMPLE SET
# =============================================================================

GAP_TRANSACTIONS = [
    ("DEPOSIT", "2020-08-03", "0", "2000", None, "Matti Meikäläinen"),
    ("FEE", "2020-08-31", "0", "-3.5", None, "Palvelumaksut 08/2020"),
    ("BUY", "2020-09-03", "40", "-652", "Aurora Metsä Oyj", "Osto Aurora Metsä Oyj (Rajakurssi)"),
    ("SELL", "2020-09-16", "-50", "644.5", "Saimaa Energia Oyj", "Myynti Saimaa Energia Oy (Päivän kurssi)"),
    ("FEE", "2020-09-30", "0", "-3.5", None, "Palvelumaksut 09/2020"),
    ("BUY", "2020-11-03", "10", "-231", "Tunturi Pankki Oyj", "Osto Tunturi Pankki Oyj (Rajakurssi)"),
    ("DEPOSIT", "2020-12-15", "0", "500", None, "Matti Meikäläinen"),
    ("FEE", "2020-12-31", "0", "-3.5", None, "Palvelumaksut 12/2020"),
    ("SELL", "2021-01-07", "-100", "428.56", "Kaamos Robotics AB", "Myynti Kaamos Robotics A (Päivän kurssi)"),
    ("SELL", "2021-01-22", "-60", "424", "Pohjola Tekniikka Oyj", "Myynti Pohjola Tekniikka (Päivän kurssi)"),
    ("DIVIDEND", "2021-02-15", "0", "28.5", "Aurora Metsä Oyj", "Aurora Metsä 5930284716 (Tuotto)"),
    ("FEE", "2021-03-01", "0", "-3.5", None, "Palvelumaksut 02/2021"),
]
# (custody file, custody line) -> cash line.
GAP_PAIRS = {(1, 2): 13, (1, 3): 12, (1, 4): 10, (2, 2): 3, (2, 3): 5, (2, 4): 4}
# cash line, row kind, zone, checkpoint: C1 = 2020-08-31 absorbs what precedes it, C2 = 2021-01-04 the trades of the gap.
GAP_CASH_ROWS = [
    pytest.param(2, "standalone", "window", "", id="fee-02-2021"),
    pytest.param(6, "standalone", "gap", "", id="gap-fee-imported"),
    pytest.param(7, "standalone", "gap", "", id="gap-deposit-imported"),
    pytest.param(8, "summarized", "gap", "2021-01-04", id="gap-sale-summarized"),
    pytest.param(9, "summarized", "gap", "2021-01-04", id="gap-purchase-summarized"),
    pytest.param(11, "standalone", "window", "", id="fee-09-2020"),
    pytest.param(14, "standalone", "before", "2020-08-31", id="fee-08-2020"),
    pytest.param(15, "standalone", "before", "2020-08-31", id="first-deposit"),
]
# The positions at C2 = 2021-01-04: E3 bounds the only movements of Kaamos (the sale of 100 on
# 2021-01-05) and of Pohjola (the sale of 60 on 2021-01-20) from below; the Tuotto of 2021-02-15 proves Aurora.
GAP_C2_POSITIONS = {"Aurora Metsä Oyj": (Decimal("190"), "exact"), "Kaamos Robotics AB": (Decimal("100"), "at_least"), "Pohjola Tekniikka Oyj": (Decimal("60"), "at_least")}


class TestGapSampleFacts:
    """The gap set: custody 2020-09-01…10-30 and 2021-01-05…02-26, one cash statement proving the trades in between."""

    def test_row_kinds(self) -> None:
        kinds = Counter(record["lf_row_kind"] for record in _records(_sample_table("gap")))

        assert kinds == Counter({"pair": 6, "standalone": 6, "summarized": 2, "excluded": 1, "truth_cash": 2, "truth_position": 4, "verification": 1})

    def test_pairs(self) -> None:
        assert _pair_lines(_records(_sample_table("gap"))) == GAP_PAIRS

    def test_a_trade_settled_after_the_statement(self) -> None:
        """The last purchase (2021-02-26) settles on 2021-03-02, after the last cash row (2021-03-01)."""
        row = _data_row(_records(_sample_table("gap")), "custody", 2, 5)

        assert (row["lf_row_kind"], row["lf_zone"], row["lf_reason"], row["lf_checkpoint"]) == ("excluded", "window", "not_yet_settled", "")

    @pytest.mark.parametrize(("line", "kind", "zone", "checkpoint"), GAP_CASH_ROWS)
    def test_cash_rows(self, line: int, kind: str, zone: str, checkpoint: str) -> None:
        row = _data_row(_records(_sample_table("gap")), "cash", 1, line)

        assert (row["lf_row_kind"], row["lf_zone"], row["lf_reason"], row["lf_checkpoint"]) == (kind, zone, "", checkpoint)

    def test_truth_rows(self) -> None:
        records = _records(_sample_table("gap"))

        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"], row["lf_zone"]) for row in _truth(records, "truth_cash")] == [("2020-08-31", "3519.86", "EUR", "before"), ("2021-01-04", "3482.36", "EUR", "gap")]
        assert [(row["custody:Sijoituskohde"], row["lf_value"], row["lf_proof"], row["lf_zone"]) for row in _truth(records, "truth_position", "2020-08-31")] == [("Saimaa Energia Oyj", "50", "at_least:E3", "before")]
        at_c2 = sorted((row["custody:Sijoituskohde"], row["lf_value"], row["lf_proof"], row["lf_zone"]) for row in _truth(records, "truth_position", "2021-01-04"))
        assert at_c2 == [("Aurora Metsä Oyj", "190", "exact:E1", "gap"), ("Kaamos Robotics AB", "100", "at_least:E3", "gap"), ("Pohjola Tekniikka Oyj", "60", "at_least:E3", "gap")]
        assert [(row["lf_checkpoint"], row["lf_value"], row["lf_currency"]) for row in _truth(records, "verification")] == [("2021-02-26", "4363.42", "EUR")]

    def test_transactions(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("gap", tmp_path_factory)

        assert _transactions(output) == [_expected_tx(*row) for row in GAP_TRANSACTIONS]

    def test_assets(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("gap", tmp_path_factory)

        assert sorted(_names(output).values()) == ["Aurora Metsä Oyj", "Kaamos Robotics AB", "Pohjola Tekniikka Oyj", "Saimaa Energia Oyj", "Tunturi Pankki Oyj"]

    def test_todos(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("gap", tmp_path_factory)

        todos = _sample_todos(output)

        assert todos == _expected_todos(GAP_TRANSACTIONS)
        assert [todo[4] for todo in todos if todo[0] in (3, 8)] == ["644.50", "428.56"]

    def test_notices(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        _table, _path, output = _sample_parse("gap", tmp_path_factory)

        assert sorted(set(_codes(output.warnings))) == ["deposit_assumed", "excluded_not_yet_settled"]

    def test_checkpoints_and_verification(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """The opening absorbs the two rows before 2020-09-01 over an opening balance of 1523.36; the gap checkpoint the two trades of the gap."""
        _table, _path, output = _sample_parse("gap", tmp_path_factory)

        opening, gap = _checkpoint(output, "2020-08-31"), _checkpoint(output, "2021-01-04")
        assert (opening.kind, _money(opening.cash), _positions(output, opening)) == ("opening", {"EUR": Decimal("3519.86")}, {"Saimaa Energia Oyj": (Decimal("50"), "at_least")})
        assert sorted((row.as_of, row.currency, row.amount, row.label) for row in opening.absorbed.rows) == [(_d("2020-08-03"), "EUR", Decimal("2000"), "Matti Meikäläinen"), (_d("2020-08-31"), "EUR", Decimal("-3.5"), "Palvelumaksut 08/2020")]
        assert (opening.absorbed.count, _money(opening.absorbed.cash), _money(opening.absorbed.opening_cash)) == (2, {"EUR": Decimal("1996.50")}, {"EUR": Decimal("1523.36")})
        assert (gap.kind, _money(gap.cash)) == ("gap", {"EUR": Decimal("3482.36")})
        assert _positions(output, gap) == GAP_C2_POSITIONS
        assert sorted((row.as_of, row.currency, row.amount, row.label) for row in gap.absorbed.rows) == [(_d("2020-11-18"), "EUR", Decimal("-515"), "Osto Aurora Metsä Oyj"), (_d("2020-12-10"), "EUR", Decimal("223"), "Myynti Tunturi Pankki Oy")]
        assert (gap.absorbed.count, _money(gap.absorbed.cash), gap.absorbed.opening_cash) == (2, {"EUR": Decimal("-292")}, [])
        assert [(item.as_of, _money(item.cash)) for item in output.verifications] == [(_d("2021-02-26"), {"EUR": Decimal("4363.42")})]


# =============================================================================
# 15. FIXTURE GUARDS — the writers of this file, not the product
# =============================================================================


def _chain_breaks(path: Path) -> List[date]:
    """Days on which the running balance of a statement does not hold, read like B0 reads it (posted rows only)."""
    posted = []
    for line in reversed([line for line in path.read_bytes().decode("latin-1").split("\n")[1:] if line]):
        day_text, label, amount_text, saldo_text, status, _check = line.split(";")
        try:
            day, amount = _from_fi_date(day_text), _number(amount_text)
        except (ValueError, InvalidOperation):
            continue
        if status == "Toteutunut" and not label.startswith("Varaus"):
            posted.append((day, amount, _number(saldo_text)))
    if not posted:
        return []
    balance = posted[0][2] - posted[0][1]
    breaks = []
    for day in sorted({row[0] for row in posted}):
        todays = [row for row in posted if row[0] == day]
        balance += sum(amount for _day, amount, _saldo in todays)
        if balance not in {saldo for _day, _amount, saldo in todays}:
            breaks.append(day)
            balance = todays[-1][2]  # carry on from what the bank states, to report each break once
    return breaks


class TestFixtureGuards:
    """Fixture guards (pass before and after phase B): the writers produce what the tests say they produce."""

    @pytest.mark.parametrize("name", ["danske_bank-cash.csv", "danske_bank-gap-cash.csv"])
    def test_fixture_guard_the_cash_writer_reproduces_the_samples(self, name: str, tmp_path: Path) -> None:
        """From its own rows, the writer rebuilds a sample byte for byte: Latin-1, ``;``, LF, newest first, the running Saldo."""
        lines = [line for line in (SAMPLE_DIR / name).read_bytes().decode("latin-1").split("\n")[1:] if line]
        rows = [_Cash(_from_fi_date(cells[0]).isoformat(), cells[1], str(_number(cells[2])), status=cells[4]) for cells in (line.split(";") for line in reversed(lines))]
        oldest = lines[-1].split(";")
        path = tmp_path / name

        _write_cash(path, rows, opening=str(_number(oldest[3]) - _number(oldest[2])))

        assert path.read_bytes() == (SAMPLE_DIR / name).read_bytes()

    def test_fixture_guard_the_custody_writer_writes_the_sample_cells(self, tmp_path: Path) -> None:
        """Text dates, numbers as cells, the unnamed column H and the account, as in the main custody sample."""
        header, rows = _sample_custody()
        path = tmp_path / "custody.xlsx"

        _write_custody(path, [_trade("2020-02-10", "Aurora Metsä Oyj", 50, 14.2, -718, value="2020-02-12"), _tuotto("2020-04-15", "Aurora Metsä Oyj", 150, 22.5), _demerger("2020-06-01", "Lumi Holding Oyj", -100, old=True)])

        written = [list(row) for row in openpyxl.load_workbook(path).active.iter_rows(values_only=True)]
        assert written == [header, rows[0], rows[1], rows[5]]

    def test_fixture_guard_a_preamble_puts_the_header_on_row_four(self, tmp_path: Path) -> None:
        path = tmp_path / "custody.xlsx"

        _write_custody(path, [_kaamos("2021-03-01")], preamble=[["Sijoitukset"], ["Tapahtumat 01.03.2021 - 31.03.2021"], [None]])

        written = [list(row) for row in openpyxl.load_workbook(path).active.iter_rows(values_only=True)]
        assert written[3] == list(CUSTODY_HEADER)

    def test_fixture_guard_scenario_statements_keep_their_chain(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Every synthetic statement holds its running balance, except the one built to break it (on 2021-03-15)."""
        breaks = {(name, path.name): _chain_breaks(path) for name in sorted(SCENARIOS) for path in _scenario_files(name, tmp_path_factory).cash}

        assert {key: days for key, days in breaks.items() if days} == {("chain", "cash-1.csv"): [_d("2021-03-15")]}
        assert _chain_breaks(SAMPLE_DIR / "danske_bank-cash.csv") == [] and _chain_breaks(SAMPLE_DIR / "danske_bank-gap-cash.csv") == []

    def test_fixture_guard_keyed_rows_land_where_the_tests_look(self, tmp_path_factory: pytest.TempPathFactory) -> None:
        """Newest first: the Saimaa row written after the Aurora row lands above it, so file order alone would cross the pairs."""
        builder = _scenario_files("by_name", tmp_path_factory)

        assert (builder.where["a"], builder.where["s"]) == (("custody", 1, 3), ("custody", 1, 4))
        assert (builder.where["cs"], builder.where["ca"]) == (("cash", 1, 3), ("cash", 1, 4))
        assert builder.row_counts == {("custody", 1): 4, ("cash", 1): 5}

    def test_fixture_guard_custody_variants_change_one_thing(self) -> None:
        """Each variant of the main custody sample changes only the header cell or the columns it names."""
        header, rows = _sample_custody()
        expected = {"br": ([5], []), "space": ([5], []), "valuutta": ([7], []), "no_currency": ([], [7]), "date_cells": ([], [0, 1]), "text_numbers_nbsp": ([], [3, 4, 5, 6]), "text_numbers_space": ([], [3, 4, 5, 6])}

        changes = {}
        for variant in CUSTODY_VARIANTS:
            changed_header, changed_rows = _custody_variant(variant)
            in_header = [index for index, (old, new) in enumerate(zip(header, changed_header, strict=True)) if old != new]
            in_rows = sorted({index for old_row, new_row in zip(rows, changed_rows, strict=True) for index, (old, new) in enumerate(zip(old_row, new_row, strict=True)) if old != new})
            changes[variant] = (in_header, in_rows)

        assert changes == expected

    def test_fixture_guard_dates_and_numbers(self) -> None:
        """The calendar the scenarios reason about (five business days, no holidays) and the Finnish number formats."""
        assert _d("2021-02-28").weekday() == 6
        assert [_business_days_after(_d(day), 5) for day in ("2021-02-28", "2021-03-31", "2021-05-31", "2021-06-30", "2021-03-09", "2020-02-02", "2020-06-26", "2020-10-30", "2021-01-04")] == [
            _d(day) for day in ("2021-03-05", "2021-04-07", "2021-06-07", "2021-07-07", "2021-03-16", "2020-02-07", "2020-07-03", "2020-11-06", "2021-01-11")
        ]
        assert [_fi_amount(Decimal(value)) for value in ("-718.00", "22.50", "-1181.04", "0.42", "6000", "0")] == ["-718", "22,5", "-1181,04", "0,42", "6000", "0"]
        assert _fi_amount(Decimal("-1181.04"), thousands="\u00a0") == "-1\u00a0181,04"
        assert [_number(value) for value in ("-1181.04", "\u22121\u00a0181,04", "22,5", "1 000")] == [Decimal("-1181.04"), Decimal("-1181.04"), Decimal("22.5"), Decimal("1000")]
        assert (_fi_text(-1181.04, thousands="\u00a0", minus="\u2212"), _fi_text(41.2, thousands=" ", minus="-"), _fi_text("EUR", thousands=" ", minus="-")) == ("\u22121\u00a0181,04", "41,2", "EUR")
